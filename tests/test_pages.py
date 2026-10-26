"""Render every page against a real database.

Two production bugs got through a green suite because nothing here ever rendered a template:
a window tuple changed shape and `/changes` raised on every request, and a mid-ingest snapshot
with no rows made the overview blow up. Unit tests cannot catch either — only rendering can.
"""

from __future__ import annotations

import html
import json
import re
from datetime import datetime, timedelta
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

from ota_analytics import config, db, ingest, registry, rollup, scheduler, sources  # noqa: E402


def api_device(imei: str, firmware: str, *, base="7.5.0.27", target=None, model="LOCAT140VB",
               task=None, ping_hours_ago=2.0) -> dict:
    when = datetime.now() - timedelta(hours=ping_hours_ago)
    return {
        "deviceId": imei, "deviceName": imei, "createdBy": 1, "createdByName": "riya",
        "creationTime": 1720009238354,
        "lastPingTime": int(when.timestamp() * 1000),
        "currFirmVer": firmware, "baseFirm": base,
        "updateFirmVer": target or firmware,
        "currConfigVersion": "2.2.2", "model": model, "hwVer": "1.2.0",
        "vin": "MAT562014RKP83714", "iccid": "8991922305932268741F",
        "type_Task": {} if task is None else task,
        "groupNames": "49A 7k, 51A 4K",
    }


@pytest.fixture
def client(tmp_path, monkeypatch):
    """A dashboard backed by a temp database holding two snapshots and a real change."""
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(config, "EXPORT_DIR", tmp_path / "exports")
    monkeypatch.setattr(scheduler, "STATE_PATH", tmp_path / "scheduler.json")
    monkeypatch.setattr(sources, "SETTINGS_PATH", tmp_path / "connection.json")
    monkeypatch.setattr(scheduler, "_scheduler", None)

    conn = db.connect()
    first = ingest.ingest_records(conn, [
        api_device("111", "7.5.0.27", target="7.5.0.51A", task={"1": 1400}),
        api_device("222", "7.5.0.51A"),
        api_device("333", "2.0.0125", model="TML_Ax1", base="2.0.0125", target="2.0.0162",
                   task={"1": 1400}, ping_hours_ago=900),
        api_device("444", "7.5.0.51A", ping_hours_ago=0.2),
    ], source_name="API test", snapshot_at=datetime.now() - timedelta(hours=6))

    second = ingest.ingest_records(conn, [
        api_device("111", "7.5.0.51A", target="7.5.0.51A"),        # completed an upgrade
        api_device("222", "7.5.0.27", base="7.5.0.27", target="7.5.0.27"),   # fell back to base
        api_device("333", "2.0.0125", model="TML_Ax1", base="2.0.0125", target="2.0.0162",
                   task={"1": 1400}, ping_hours_ago=900),
        api_device("444", "7.5.0.51A", ping_hours_ago=0.2),
    ], source_name="API test", snapshot_at=datetime.now())

    for result in (first, second):
        rollup.rollup_snapshot(conn, result.snapshot_id)

    from ota_analytics import api
    return TestClient(api.app, raise_server_exceptions=False)


PAGES = ["/", "/pending", "/firmware", "/changes", "/devices", "/reachability",
         "/groups", "/quality", "/update", "/errors", "/web-cota", "/cota", "/cota/devices",
         "/cota/signin", "/cota/jobs/new", "/cota/commands"]


@pytest.mark.parametrize("path, module, page", [
    ("/", "Web FOTA", "Overview"), ("/devices", "Web FOTA", "Devices"),
    ("/update", "Web FOTA", "Update data"), ("/web-cota", "Web COTA", "Configuration"),
    ("/cota", "Intouch COTA", "Jobs"), ("/cota/devices", "Intouch COTA", "Devices"),
    ("/cota/jobs/new", "Intouch COTA", "Jobs"), ("/cota/commands", "Intouch COTA", "Commands"),
])
def test_the_rail_lists_every_module_and_marks_the_current_page(client, path, module, page):
    body = client.get(path).text
    groups = re.findall(r'<div class="rail-group[^"]*">(.*?)</div>', body)
    assert groups == ["Web FOTA", "Web COTA", "Intouch COTA"]
    on = re.findall(r'class="rail-item on[^"]*"[^>]*>.*?<span class="rail-label">(.*?)</span>',
                    body, re.S)
    assert on == [page], f"{path} marks {on} as current"
    # The top bar names the page and the module it belongs to.
    assert f"<h1>{page}</h1>" in body
    assert f"<strong>{module}</strong>" in body


def test_cota_sign_in_is_reached_from_its_status_chip(client, cota_env):
    """No separate Sign in button: the "Cloud: …" chip says the state and opens the page."""
    body = client.get("/cota/signin").text
    rail = body[body.index('<aside class="rail"'):body.index("</aside>")]
    assert "/cota/signin" not in rail
    assert 'id="cota-signin-btn"' not in body
    assert not re.search(r">\s*Cloud sign-in\s*</(a|button)>", body)    # no button; the card heading may say it
    assert re.search(r'class="status-chip status-\w+ is-current"\s+href="/cota/signin"', body)
    assert "<h1>Sign in</h1>" in body and "<strong>Intouch COTA</strong>" in body
    jobs = client.get("/cota").text
    assert 'id="cota-signin-chip"' in jobs and "is-current" not in jobs.split('id="cota-signin-chip"')[0][-200:]
    assert 'id="cota-signin-chip"' not in client.get("/").text

def test_web_cota_is_listed_as_not_built(client):
    body = client.get("/").text
    item = re.search(r'<a href="/web-cota"[^>]*>.*?</a>', body, re.S).group(0)
    assert "is-soon" in item and ">Soon<" in item


def test_the_cota_modules_do_not_carry_fotas_header(client, cota_env):
    """The fetch chips and Update data button belong to the snapshot warehouse. On a COTA page
    they would claim a freshness that has nothing to do with what is on screen."""
    for path in ("/cota", "/cota/devices", "/cota/signin", "/web-cota"):
        body = client.get(path).text
        assert 'id="agent-chip"' not in body
        assert 'href="/update">Update data</a>' not in body
    assert 'id="agent-chip"' in client.get("/").text
    assert 'href="/cota/signin"' in client.get("/cota").text      # COTA's own chip instead


class FakeKeyring:
    """Stands in for Windows Credential Manager, so no test writes to the real one."""

    def __init__(self):
        self.store = {}

    def set_password(self, service, user, secret):
        self.store[(service, user)] = secret

    def get_password(self, service, user):
        return self.store.get((service, user))

    def delete_password(self, service, user):
        self.store.pop((service, user), None)


@pytest.fixture(autouse=True)
def cota_env(tmp_path, monkeypatch):
    """Every test here, not only the COTA ones: pages read the sign-in state on render, and
    without this they would read — and the sign-in tests write — the real credential store."""
    from ota_analytics import cota, cota_connection

    keyring = FakeKeyring()
    monkeypatch.setattr(sources, "_keyring", lambda: keyring)
    # Jobs must finish the day they start, and these tests run on the real clock: pinned to ten
    # hours left, a few-minute test job would otherwise be refused just before midnight.
    from ota_analytics import cota_campaign
    monkeypatch.setattr(cota_campaign, "minutes_left_today", lambda now=None: 600.0)
    monkeypatch.setattr(cota_connection, "SETTINGS_PATH", tmp_path / "cota_connection.json")
    monkeypatch.delenv(cota.ENV_TOKEN, raising=False)
    monkeypatch.delenv(cota_connection.ENV_PASSWORD, raising=False)
    return keyring


def test_a_pasted_token_is_kept_but_never_shown(client, cota_env):
    secret = "3f7c2a10-not-a-real-token-but-must-never-render"
    assert "Cloud: not signed in" in client.get("/cota").text

    body = client.post("/cota/signin", data={"cloud": "intouch", "method": "token",
                                             "token": f"Bearer {secret}"}).text
    assert "Token saved" in body
    assert secret not in body
    assert secret in cota_env.store.values()          # "Bearer " stripped, kept in the store
    assert "Cloud: signed in" in client.get("/cota").text
    for path in ("/cota", "/cota/devices", "/cota/signin"):
        assert secret not in client.get(path).text

    from ota_analytics import cota_connection
    assert secret not in cota_connection.SETTINGS_PATH.read_text(encoding="utf-8")

    client.post("/cota/signout")
    assert not cota_env.store
    assert "Cloud: not signed in" in client.get("/cota").text


def test_signing_out_keeps_a_remembered_password_and_forgetting_is_its_own_step(client, cota_env):
    """10-10-2026: sign out, and the next sign-in still needs no typing; forgetting the password
    is a red button that asks first."""
    from ota_analytics import cota_connection
    conn = cota_connection.load()
    conn.method, conn.username = "password", "riya@example.com"
    cota_connection.save(conn)
    cota_connection.save_password("riya@example.com", "not-a-real-password")
    cota_connection.use_token(conn, "a-token-long-enough-to-count")
    body = client.post("/cota/signout").text
    assert "Your password stays saved" in body and "not-a-real-password" not in body
    assert "not-a-real-password" in cota_env.store.values()           # still in the store
    assert "Cloud: not signed in" in client.get("/cota").text
    page = re.sub(r"\s+", " ", client.get("/cota/signin").text)
    assert 'placeholder="Saved"' in page and 'form="cota-forget"' in page
    assert 'data-confirm="Forget the saved password for {name}?' in page
    body = client.post("/cota/signin/forget").text
    assert "The saved password was removed." in body and not cota_env.store


def test_the_sign_in_page_keeps_only_what_is_needed(client, cota_env):
    page = re.sub(r"\s+", " ", client.get("/cota/signin").text)
    assert ">CVTMS Cloud</option>" in page and "InTouch cloud" not in page
    assert "Tokens and passwords are kept in" not in page
    assert "Nothing is sent to a device from this tab" not in page
    assert "The cloud cannot tell us when a token expires" not in page
    assert page.index("<summary") < page.index('name="base_url"')     # API URL is in the fold


def test_a_token_from_the_environment_is_reported_as_such(client, cota_env, monkeypatch):
    from ota_analytics import cota

    monkeypatch.setenv(cota.ENV_TOKEN, "eyJ-from-the-environment")
    body = client.get("/cota/signin").text
    assert "Cloud: token from environment" in body
    assert "eyJ-from-the-environment" not in body


def test_password_sign_in_on_a_custom_cloud_without_a_login_url_says_why(client, cota_env):
    """A cloud this tool does not know has no login address to fall back on, and guessing one
    would send a password somewhere nobody verified."""
    body = client.post("/cota/signin", data={"cloud": "custom", "base_url": "https://x.example/api",
                                             "method": "password", "username": "ops",
                                             "password": "hunter2-never-echoed"}).text
    assert "No login URL" in body
    assert "hunter2-never-echoed" not in body
    assert not cota_env.store


def _login_cloud(monkeypatch, *, token="0f3c9e2a-1b2c-4d5e-8f90-a1b2c3d4e5f6", sends=None):
    """A fake CTVMS: /user/login hands out `token`; the COTA calls accept only the newest one."""
    import httpx

    state = {"token": token, "logins": [], "sends": sends if sends is not None else []}

    def handler(request):
        if request.url.path.endswith("/login"):
            state["logins"].append(request)
            return httpx.Response(200, json={"status": True, "data": {"token": state["token"]}})
        if request.headers.get("authorization") != f"Bearer {state['token']}":
            return httpx.Response(401, json={"message": "Unauthorized"})
        if request.url.path.endswith("/saveCOTAConfig/0"):
            state["sends"].append(json.loads(request.content))
            return httpx.Response(200, json={"msg": "Command send Successfully.."})
        return httpx.Response(200, json={"data": []})

    real = httpx.Client
    monkeypatch.setattr(httpx, "Client",
                        lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    return state


def test_password_sign_in_uses_the_portals_own_login_request(client, cota_env, monkeypatch):
    """Captured from the CTVMS portal (2026-10-06): POST IntouchAdminApi/user/login, multipart,
    username + MD5 of the password."""
    import hashlib

    from ota_analytics import cota_connection

    state = _login_cloud(monkeypatch)
    body = client.post("/cota/signin", data={"cloud": "intouch", "method": "password",
                                             "username": "ops", "password": "pw-never-echoed",
                                             "remember": "true"}).text
    assert "Signed in as ops" in body and "pw-never-echoed" not in body
    request = state["logins"][0]
    assert str(request.url) == "https://ctvms.mappls.com/IntouchAdminApi/user/login"
    assert request.headers["content-type"].startswith("multipart/form-data")
    md5 = hashlib.md5(b"pw-never-echoed").hexdigest()
    assert b'name="username"\r\n\r\nops' in request.content
    assert f'name="password"\r\n\r\n{md5}'.encode() in request.content
    assert b"pw-never-echoed" not in request.content            # never the plain password
    assert state["token"] in cota_env.store.values()
    assert cota_env.store[(sources.SERVICE_NAME, "cota-password:ops")] == "pw-never-echoed"
    assert sources.load_password("ops") is None                  # not FOTA's account
    assert cota_connection.load().login_url == "https://ctvms.mappls.com/IntouchAdminApi/user/login"


def test_the_preset_login_cannot_be_redirected_by_a_posted_url(client, cota_env, monkeypatch):
    from ota_analytics import cota_connection

    state = _login_cloud(monkeypatch)
    client.post("/cota/signin", data={"cloud": "intouch", "method": "password", "username": "ops",
                                      "password": "pw", "login_url": "https://evil.example/login",
                                      "login_encoding": "json", "password_hash": "none"})
    assert str(state["logins"][0].url).startswith("https://ctvms.mappls.com/")
    conn = cota_connection.load()
    assert (conn.login_url, conn.login_encoding, conn.password_hash) == (
        "https://ctvms.mappls.com/IntouchAdminApi/user/login", "multipart", "md5")
    assert 'readonly' in client.get("/cota/signin").text.split('name="login_url"', 1)[1][:200]


def test_a_custom_cloud_signs_in_where_it_is_told(client, cota_env, monkeypatch):
    state = _login_cloud(monkeypatch)
    body = client.post("/cota/signin", data={
        "cloud": "custom", "base_url": "https://other.example/api", "method": "password",
        "username": "ops", "password": "pw", "login_url": "https://other.example/login",
        "login_encoding": "json"}).text
    assert "Signed in as ops" in body
    assert str(state["logins"][0].url) == "https://other.example/login"
    assert b'"username"' in state["logins"][0].content


def _remembered_password_sign_in(client, monkeypatch):
    state = _login_cloud(monkeypatch, token="token-one-0123456789")
    client.post("/cota/signin", data={"cloud": "intouch", "method": "password", "username": "ops",
                                      "password": "pw", "remember": "true"})
    state["token"] = "token-two-0123456789"          # the cloud expires the first token
    return state


def test_an_expired_token_renews_itself_on_a_check(client, cota_env, monkeypatch):
    state = _remembered_password_sign_in(client, monkeypatch)
    out = client.post("/cota/console/check", data={"device": "14906"}).json()
    assert out["ok"] and "signed in again" in out["message"]
    assert out["cloud"]["label"] == "Cloud: signed in"
    assert len(state["logins"]) == 2 and "token-two-0123456789" in cota_env.store.values()


def test_an_expired_token_renews_itself_and_the_send_goes_once(client, cota_env, monkeypatch):
    from ota_analytics import db

    state = _remembered_password_sign_in(client, monkeypatch)
    reply = client.post("/cota/console/send", data={"device_id": "14906", "device_type": "124",
                                                    "cmd_type": "36", "val1": "ONCE"},
                        follow_redirects=False)
    assert reply.status_code == 303
    assert [b["val1"] for b in state["sends"]] == ["ONCE"]     # reached the cloud exactly once
    states = [r[0] for r in db.connect().execute("SELECT state FROM cota_task ORDER BY id")]
    assert states == ["send_failed", "sent"]                    # the refusal stays on the record
    assert "Cloud: signed in" in client.get(reply.headers["location"]).text


def test_without_a_saved_password_an_expired_token_stays_expired(client, cota_env, monkeypatch):
    from ota_analytics import cota

    state = _login_cloud(monkeypatch, token="token-two-0123456789")
    cota.save_token("token-one-0123456789")
    out = client.post("/cota/console/check", data={"device": "14906"}).json()
    assert not out["ok"] and out["cloud"]["label"] == "Cloud: session expired"
    assert state["logins"] == []


def test_a_custom_cloud_keeps_its_url_and_a_known_one_cannot_be_edited(client, cota_env):
    from ota_analytics import cota, cota_connection

    client.post("/cota/signin", data={"cloud": "custom", "base_url": "https://other.example/api/",
                                      "method": "token", "token": "t"})
    assert cota_connection.load().base_url == "https://other.example/api"
    client.post("/cota/signin", data={"cloud": "intouch", "base_url": "https://evil.example",
                                      "method": "token", "token": "t"})
    assert cota_connection.load().base_url == cota.BASE_URL


def test_the_cota_page_lists_jobs_from_this_install(client, cota_env):
    from ota_analytics import db

    conn = db.connect()
    with conn:
        conn.execute("INSERT INTO cota_job (id, name, created_at) "
                     "VALUES (7, 'BLE MAC rollout', '2026-10-05 12:00:00')")
    body = client.get("/cota").text
    assert "BLE MAC rollout" in body
    assert "No jobs yet" not in body


def test_the_model_picker_is_the_same_control_on_every_page(client):
    """Overview and Firmware both slice by model, so they must not do it two different ways.

    Firmware used an always-open multi-select list box that needed "ctrl+click to pick several"
    written underneath it; the overview had a dropdown. Same job, same control.
    """
    for path in ("/", "/firmware"):
        body = client.get(path).text
        # The Devices tables' list filter (12-10-2026): every model with its count, applied
        # through the page's own form.
        assert 'id="model-form"' in body, f"{path} has no model filter"
        assert 'data-multi="model" data-form="model-form"' in body
        assert 'class="multi-count"' in body
        assert "ctrl+click" not in body, f"{path} still explains a multi-select"
        assert "<select name=\"model\" multiple" not in body


def test_the_firmware_picker_keeps_the_snapshot_being_viewed(client):
    """Narrowing by model must not silently jump back to the latest snapshot."""
    from ota_analytics import db, metrics

    oldest = metrics.snapshots(db.connect())[-1]["id"]
    body = client.get(f"/firmware?snapshot={oldest}").text
    assert f'name="snapshot" value="{oldest}"' in body


def test_selecting_a_model_on_the_firmware_page_filters_it(client):
    from ota_analytics import db, metrics

    conn = db.connect()
    models = [r["label"] for r in metrics.task_state_by(conn, metrics.latest_snapshot_id(conn),
                                                        "model")]
    assert len(models) > 1, "fixture needs at least two models to prove filtering"

    body = client.get(f"/firmware?model={models[0]}").text
    assert body.count(f'value="{models[0]}"') >= 1

    # The summary reports the narrowing rather than still claiming the whole fleet. Compared
    # with whitespace collapsed, because the template wraps the line and the exact run of
    # spaces is not the thing under test.
    import re
    flat = re.sub(r"\s+", " ", body)
    assert f"in 1 of {len(models)} models" in flat
    assert "all {} models".format(len(models)) not in flat


def test_the_firmware_table_names_each_denominator_and_task_column(client):
    """Three percentage columns sit side by side and measure against different totals.

    One repeated "Share (%)" would put the same word on three different meanings in adjacent
    columns. The group row carries the subject, the row under it carries the denominator — and
    now a Task figure for each, so "how many of these are still waiting" is answerable per
    group rather than only for the row as a whole.
    """
    import re

    body = client.get("/firmware").text
    flat = re.sub(r"\s+", " ", body)

    assert "% of fleet" in flat and "% of version" in flat
    assert flat.count("% of version") == 2          # online and offline, not the fleet share
    assert "Share (%)" not in flat                  # the ambiguous label it replaced

    # Read the group row itself rather than searching the whole page, so a heading appearing
    # somewhere else cannot make this pass.
    header = flat[flat.index('<tr class="group-row">'):flat.index("</thead>")]
    for heading in ("Model", "Firmware", "Devices", "Online", "Offline", "Distribution"):
        assert f">{heading}<" in header, f"{heading} is missing from the header"

    # Three groups, each Count / % / Task.
    assert header.count("Count") == 3
    assert header.count(">Task<") == 3

    # The column of zeros is gone: Inactive read 0 on 102 of 103 rows on the real fleet.
    assert "Inactive" not in header


def test_devices_that_never_pinged_are_still_accounted_for(client):
    """Online + Offline does not reach the row total: STATUS has a third value.

    The Inactive column was dropped because it reads 0 on 102 of 103 rows on the real fleet —
    but those 645 devices must not vanish silently, or one row loses most of itself with nothing
    to explain it. The figure moved into a marker on that row instead of a column of zeros.
    """
    from ota_analytics import db, metrics

    conn = db.connect()
    rows = metrics.firmware_mix(conn, metrics.latest_snapshot_id(conn))
    assert rows, "fixture has no firmware rows"
    for row in rows:
        assert row["online"] + row["offline"] + row["inactive"] == row["devices"], row

    template = Path("ota_analytics/web/templates/firmware.html").read_text(encoding="utf-8")
    assert "never pinged" in template
    assert "footmark" in template


def test_each_group_reports_its_own_pending_count(client):
    """Task pending is split the same way the fleet is, because the split is the point.

    A task pending on a reachable device is the one worth chasing; on a dark one it is parked.
    Reported per version so a rollout stalling on one build stands out from one merely waiting
    for vehicles to be switched on.
    """
    from ota_analytics import db, metrics

    conn = db.connect()
    for row in metrics.firmware_mix(conn, metrics.latest_snapshot_id(conn)):
        assert row["pending_online"] <= row["online"]
        assert row["pending_offline"] <= row["offline"]
        # The two never exceed the row's total pending: an Inactive device can be pending too,
        # so they may sum to less, but never to more.
        assert row["pending_online"] + row["pending_offline"] <= row["pending"]


@pytest.mark.parametrize("path", PAGES)
def test_every_page_renders(client, path):
    response = client.get(path)
    assert response.status_code == 200, response.text[:400]
    assert "Something went wrong" not in response.text


@pytest.mark.parametrize("path,query", [
    ("/", "window=1h"), ("/", "window=yesterday"), ("/", "window=all"),
    ("/", "model=LOCAT140VB"), ("/", "model=LOCAT140VB&model=TML_Ax1"),
    ("/changes", "window=6h"), ("/changes", "window=month"),
    ("/devices", "sort=firmware&dir=asc"), ("/devices", "status=Online&changed=24h"),
    ("/devices", "queue_state=pending&sort=seen&dir=desc"),
    ("/firmware", "model=LOCAT140VB"),
])
def test_pages_render_with_their_filters(client, path, query):
    response = client.get(f"{path}?{query}")
    assert response.status_code == 200, response.text[:400]


def test_overview_shows_the_change_that_happened(client):
    """111 upgraded and 222 fell back to base — both must appear."""
    body = client.get("/?window=all").text
    assert "What changed" in body
    assert "Fell back" in body


def test_changes_page_lists_the_fallback(client):
    body = client.get("/changes?window=all").text
    assert "222" in body                    # the device that returned to base
    assert "7.5.0.27" in body               # its base firmware


def test_json_endpoints_respond(client):
    for path in ("/api/kpis", "/api/version", "/api/agent", "/api/pending", "/api/quality"):
        response = client.get(path)
        assert response.status_code == 200, path
        assert response.json() is not None


def test_a_snapshot_still_being_written_is_not_served(client, tmp_path):
    """Ingest writes the snapshot row before its devices; that gap broke every page."""
    conn = db.connect()
    conn.execute("INSERT INTO snapshot (source_file, file_sha256, snapshot_at, ts_source, "
                 "row_count, ingested_at) VALUES ('half done','sha-x',?,'api',0,?)",
                 (datetime.now().isoformat(sep=" ", timespec="seconds"),
                  datetime.now().isoformat(sep=" ", timespec="seconds")))
    conn.commit()

    response = client.get("/")
    assert response.status_code == 200
    assert "half done" not in response.text


def test_an_unknown_window_falls_back_instead_of_failing(client):
    assert client.get("/changes?window=nonsense").status_code == 200
    assert client.get("/?window=nonsense").status_code == 200


def test_an_unknown_sort_column_is_ignored(client):
    assert client.get("/devices?sort=DROP+TABLE&dir=sideways").status_code == 200


def test_version_is_reported_everywhere(client):
    from ota_analytics import __version__

    assert client.get("/api/version").json()["version"] == __version__
    assert f"v{__version__}" in client.get("/").text


# ─── one status chip, one theme switch, one table idiom ─────────────────────

def test_the_header_carries_one_freshness_chip_not_two(client):
    """"Updated 11:08 · 8 min ago" and "Next 4:02" answer the same question.

    Side by side they read as two unrelated clocks; together they say when the data was true and
    when it will next be checked.
    """
    import re

    flat = re.sub(r"\s+", " ", client.get("/").text)
    assert flat.count('class="status-chip') == 2      # connection state, and this one
    assert flat.count('id="agent-chip"') == 1
    assert "Updated" in flat
    # The countdown writes into the same chip rather than a second one.
    assert 'id="agent-text"' in flat
    chip = flat[flat.index('id="agent-chip"'):]
    chip = chip[:chip.index("</a>")]
    assert "Updated" in chip and 'id="agent-text"' in chip


def test_a_theme_can_be_chosen_or_left_to_the_system(client):
    """Three states, and the CSS ordering is what makes all three reachable.

    A light palette defined only inside a prefers-color-scheme query could never be turned on by
    someone whose system is dark, so each explicit choice is stated on its own.
    """
    from pathlib import Path

    body = client.get("/").text
    assert 'id="theme-btn"' in body
    # Applied before the stylesheet paints, or a reader who chose light gets a flash of dark.
    assert body.index("ota-theme") < body.index("</head>")

    css = Path("ota_analytics/web/static/app.css").read_text(encoding="utf-8")
    assert ':root[data-theme="light"]' in css
    assert ':root[data-theme="dark"]' in css
    # The system preference must not override an explicit choice.
    assert ':root:not([data-theme="dark"])' in css


def test_both_grouped_tables_share_the_same_idiom(client):
    """Firmware and the model table are read the same way, without being identical.

    The model table carries more — it has room, at six rows against a hundred — so this checks
    the idiom rather than the column count: a two-tier header, every percentage named by its
    denominator, and a totals row pinned with the header instead of left at the foot.
    """
    import re

    def header_of(path, table_id):
        flat = re.sub(r"\s+", " ", client.get(path).text)
        assert f'id="{table_id}"' in flat, f"{path} has no {table_id} table"
        block = flat[flat.index(f'id="{table_id}"'):]
        return block[:block.index("</thead>")]

    firmware = header_of("/firmware", "firmware-mix")
    models = header_of("/", "model-states")

    for header in (firmware, models):
        assert 'class="group-row"' in header and 'class="sub-row"' in header
        assert 'class="totals-row"' in header
        for group in ("Devices", "Online", "Offline"):
            assert f">{group}<" in header
        # A bare "Share" repeated across columns with different denominators is what this
        # replaced; every percentage says what it is a percentage of.
        assert "Share (%)" not in header
        assert "% of" in header

    assert "% of fleet" in firmware and "% of version" in firmware
    assert "% of fleet" in models and "% of model" in models


def test_the_model_table_adds_up_two_ways(client):
    """Two accountings of the same devices, and each has to reach the row total.

        Completed + Pending + No task  = Devices
        Online + Offline + Act-Pending = Devices

    The second is why Activation-Pending is a column here and a marker on the firmware table:
    without it the (unknown) row reads 0 online and 36 offline out of 513 and loses 477 devices
    with nothing to explain them.
    """
    from ota_analytics import db, metrics

    conn = db.connect()
    rows = metrics.task_state_by(conn, metrics.latest_snapshot_id(conn), "model")
    assert rows, "fixture has no models"

    for row in rows:
        assert row["completed"] + row["pending"] + row["never_tasked"] == row["devices"], row
        assert row["online"] + row["offline"] + row["inactive"] == row["devices"], row

    header = client.get("/").text
    for heading in ("Completed", "No task", "Activation-", "Pending"):
        assert heading in header


def test_the_model_table_colours_what_is_worth_acting_on(client):
    """Pending-while-online is the number to chase; the totals are not coloured at all.

    The yellow one is conditional on being non-zero, so it is asserted against the template —
    this fixture has no reachable pending device to render it.
    """
    from pathlib import Path as _Path

    template = _Path("ota_analytics/web/templates/overview.html").read_text(encoding="utf-8")
    assert "tone-yellow-text" in template and "if r.pending_reachable" in template
    assert "tone-orange-text" in template

    # The Devices group has no Task column: that figure is Pending, and printing it twice under
    # two headings invites the reader to look for a difference that cannot exist.
    header = re.sub(r"\s+", " ", client.get("/").text)
    block = header[header.index('id="model-states"'):]
    block = block[:block.index("</thead>")]
    assert block.count(">Task<") == 2          # Online and Offline only


def test_each_group_in_the_model_table_reports_its_own_pending(client):
    from ota_analytics import db, metrics

    conn = db.connect()
    for row in metrics.task_state_by(conn, metrics.latest_snapshot_id(conn), "model"):
        assert row["pending_reachable"] <= row["online"]
        assert row["pending_offline"] <= row["offline"]
        # They may sum to less than the row's total pending — a device that has never pinged can
        # carry a task too — but never to more.
        assert row["pending_reachable"] + row["pending_offline"] <= row["pending"]


# ─── the devices filter row ─────────────────────────────────────────────────

def test_every_column_with_known_values_filters_as_in_excel(client):
    """Intouch COTA's Devices design on Web FOTA's list (12-10-2026): a list filter under each
    such header — every value with its count, All / None / Invert / Only — and a typed filter
    for the IMEI. Several versions are usually interesting together, so the firmware filter
    takes any number of them."""
    flat = re.sub(r"\s+", " ", client.get("/devices").text)
    for key in ("model", "firmware", "changed", "target", "config", "hw", "status",
                "queue_state", "seen"):
        assert f'data-multi="{key}"' in flat, f"no list filter for {key}"
    assert 'name="q" form="filters"' in flat
    assert '<select name="firmware"' not in flat
    panel = flat[flat.index('data-multi="firmware"'):]
    panel = panel[:panel.index('class="multi-foot"')]
    assert re.findall(r'type="checkbox" value="([^"]+)"', panel), "no versions offered"
    assert 'class="multi-count"' in panel
    for word in ('data-pick="all"', 'data-pick="none"', 'data-pick="invert"', "data-only"):
        assert word in panel


def test_several_firmware_versions_can_be_selected_at_once(client):
    from ota_analytics import db, metrics

    conn = db.connect()
    sid = metrics.latest_snapshot_id(conn)
    versions = [r["firmware"] for r in metrics.firmware_mix(conn, sid) if r["firmware"]][:2]
    assert len(versions) == 2, "fixture needs two firmware versions"

    one = client.get(f"/devices?firmware={versions[0]}").text
    both = client.get(f"/devices?firmware={versions[0]}&firmware={versions[1]}").text

    def total(body):
        return int(re.search(r'class="hint">([\d,]+) ', re.sub(r"\s+", " ", body))
                   .group(1).replace(",", ""))

    assert total(both) > total(one), "adding a version did not widen the selection"
    # Both stay ticked, so the control shows what it is filtering by.
    flat = re.sub(r"\s+", " ", both).replace('checked=""', "checked")
    for version in versions:
        assert f'value="{version}" checked' in flat


def test_the_group_text_box_is_gone_but_the_link_still_works(client):
    """It was an exact-match field nobody could type from memory. The Groups page links here
    with ?group=…, so the capability stays even though the control does not."""
    flat = re.sub(r"\s+", " ", client.get("/devices").text)
    assert 'placeholder="exact name"' not in flat

    from ota_analytics import db, metrics
    conn = db.connect()
    groups = metrics.groups(conn, metrics.latest_snapshot_id(conn))
    if groups:
        name = groups[0]["group_name"]
        filtered = client.get(f"/devices?group={name}")
        assert filtered.status_code == 200
        # And the filter survives paging, so it is carried in a hidden field.
        flat_filtered = re.sub(r"\s+", " ", filtered.text)
        assert f'name="group" value="{name}"' in flat_filtered
        assert 'class="group-chip"' in flat_filtered       # and it says which group is in view


def test_an_imei_can_be_searched_by_any_part_of_it(client):
    """The last few digits are what someone reads off a label, and a paste from the platform
    arrives wrapped in quotes and commas."""
    body = client.get("/devices").text
    imeis = re.findall(r'class="mono">(\d+)<', body)
    assert imeis, "fixture has no devices"
    target = imeis[0]

    def found(term):
        import urllib.parse
        page = client.get("/devices?q=" + urllib.parse.quote(term)).text
        return re.findall(r'class="mono">(\d+)<', page)

    assert found(target) == [target]
    assert target in found(target[-2:])
    # Non-digits are stripped rather than rejected, so a paste works as-is.
    assert found(f'"{target}", ') == [target]
    # And nothing matches an IMEI that is not there.
    assert found("00000000000000") == []


# ─── Intouch COTA: sending one command from the Devices tab ────────────────



class FakeCloud:
    """Stands in for cota.Client: records every body and answers as told."""

    def __init__(self, status=200, reply='{"status":"success"}', error=None):
        self.sent, self.status, self.reply, self.error = [], status, reply, error
        self.checks, self.reply_status, self.replies_body = [], 200, '{"data": []}'

    def responses(self, device_id, start, end):
        self.checks.append((device_id, start, end))
        return self.reply_status, self.replies_body

    def send(self, payload):
        if self.error:
            raise self.error
        self.sent.append(payload)
        return self.status, self.reply

    def close(self):
        pass


@pytest.fixture
def cloud(monkeypatch, cota_env):
    from ota_analytics import cota, cota_connection

    fake = FakeCloud()
    monkeypatch.setattr(cota_connection, "client", lambda conn=None: fake)
    monkeypatch.setattr(cota, "CALL_INTERVAL_SECONDS", 0)
    cota.save_token("a-token-long-enough-to-count")
    return fake


# ─── Intouch COTA: Configure — one device, one command, and its stages ─────

def _console_send(client, **over):
    form = {"device_id": "14906", "device_type": "124", "cmd_type": "36",
            "val1": "DBD76B82D531"}
    form.update(over)
    return client.post("/cota/console/send", data=form, follow_redirects=False)


def _map(conn, imei="865510083360422", device_id=14906, device_type=124):
    from ota_analytics import cota

    with conn:
        cota.put_device(conn, device_id, imei=imei, device_type=device_type, source="map.csv")


def cloud_record(timestamp=None, **over):
    """The getGPRSCommand record exactly as captured from the cloud (2026-10-06)."""
    record = {"id": "865510083360422_6AC48E16", "deviceId": 14906,
              "timestamp": timestamp or int(datetime.now().timestamp()), "type": 36,
              "val1": "245A454ED731323334D7DAD76F4BD76AC48E16D73542D1", "val2": None,
              "movementState": None, "status": 0, "commandType": None, "commandValue": None,
              "response": None, "responseTime": None, "inputVal1": None, "inputVal2": None,
              "commandId": None, "imei": "865510083360422", "commandExcecuteType": 0}
    record.update(over)
    return record


def _thread_of(body: str) -> str:
    """Just the conversation, without the page around it."""
    return body.split('id="thread"', 1)[1].split('id="composer"', 1)[0]


def _check(client, cloud, *records):
    cloud.replies_body = json.dumps({"data": list(records)})
    return client.post("/cota/console/check", data={"device": "14906"}).json()


def test_configure_is_the_first_cota_page(client, cota_env):
    body = client.get("/cota/console").text
    on = re.findall(r'class="rail-item on[^"]*"[^>]*>.*?<span class="rail-label">(.*?)</span>',
                    body, re.S)
    assert on == ["Configure"] and "<h1>Configure</h1>" in body
    rail = body[body.index('<aside class="rail"'):body.index("</aside>")]
    assert rail.index("/cota/console") < rail.index('href="/cota"')
    assert "Start configuration" in body and "No devices yet" in body


def test_model_124_and_type_36_are_set_from_edit_beside_the_name(client, cota_env):
    body = client.get("/cota/console?device=14906").text
    head = body.split('class="thread-title-row"', 1)[1].split('class="thread-counts"', 1)[0]
    assert re.search(r'name="device_type" form="composer" class="mono"\s+inputmode="numeric" value="124"', head)
    assert re.search(r'name="cmd_type" form="composer" class="mono"\s+inputmode="numeric" value="36"', head)
    assert ">Edit</summary>" in head and "Zenithra Command" in head
    composer = body.split('id="composer"', 1)[1].split("</form>", 1)[0]
    assert 'name="device_type"' not in composer and "Zenithra Command</span>" not in composer
    assert "fixed-row" not in composer
    start = client.get("/cota/console").text
    assert re.search(r'name="type" inputmode="numeric"\s+class="mono" value="124"\s+readonly', start)


def test_starting_with_an_imei_from_the_map_opens_its_device(client, cota_env):
    from ota_analytics import db

    _map(db.connect())
    body = client.get("/cota/console?device=865510083360422").text
    assert 'id="thread-title">ID: 14906 | IMEI: 865510083360422<' in body


def test_an_imei_the_map_does_not_know_is_not_guessed(client, cota_env):
    body = client.get("/cota/console?device=865510083369999").text
    assert "is not in the device map" in body
    assert 'id="composer"' not in body


def test_a_send_shows_one_tick_and_a_refresh_cannot_resend(client, cloud):
    reply = _console_send(client)
    assert reply.status_code == 303
    assert reply.headers["location"].startswith("/cota/console?device=14906&type=124")
    for _ in range(2):                                  # load it, then "refresh"
        body = client.get(reply.headers["location"]).text
    assert cloud.sent == [{"deviceType": [124], "deviceList": [14906], "type": 36,
                           "val1": "DBD76B82D531"}]
    assert "Zenithra Command · 36" in body and "DBD76B82D531" in body
    in_thread = _thread_of(body)
    assert 'class="tick tick-accepted"' in in_thread and 'class="tick tick-api"' not in in_thread
    assert 'class="waiting"' in body                    # watching for the cloud's record
    assert 'id="thread-title">ID: 14906<' in body       # no IMEI known yet


def test_the_clouds_record_gives_two_ticks_the_imei_and_the_full_preview(client, cloud):
    from ota_analytics import db

    _console_send(client)
    out = _check(client, cloud, cloud_record())
    assert out["ok"] and out["latest_stage"] == "api"
    assert out["label"] == "ID: 14906 | IMEI: 865510083360422"
    page = html.unescape(out["html"])
    assert 'class="tick tick-api"' in page
    assert 'data-for="the device"' in page             # the cloud has it; now the device
    assert "Held as" not in page                       # no "response via API" bubble any more
    assert "865510083360422_6AC48E16" in page and "Raw API response" in page
    assert re.search(r"status\s+0 — waiting for the device", page)
    # The preview shows what was typed AND what the cloud holds, rewritten val1 and all.
    assert '"val1": "DBD76B82D531"' in page
    assert '"val1": "245A454ED731323334D7DAD76F4BD76AC48E16D73542D1"' in page
    assert '"commandExcecuteType": 0' in page

    conn = db.connect()
    assert tuple(conn.execute("SELECT imei, device_id, device_type, source FROM cota_device")
                 .fetchone()) == ("865510083360422", 14906, 124, "cloud record")
    assert conn.execute("SELECT imei FROM cota_task").fetchone()[0] == "865510083360422"
    matched = conn.execute("SELECT task_id IS NOT NULL FROM cota_command").fetchone()[0]
    assert matched == 1


def test_checking_again_never_duplicates_a_command(client, cloud):
    from ota_analytics import db

    _console_send(client)
    _check(client, cloud, cloud_record())
    out = _check(client, cloud, cloud_record(status=1))
    assert db.connect().execute("SELECT COUNT(*), MAX(status) FROM cota_command").fetchone()[:] \
        == (1, 1)
    assert html.unescape(out["html"]).count("865510083360422_6AC48E16 ") <= 2


def test_a_record_far_from_any_send_is_shown_as_sent_elsewhere(client, cloud):
    hour_ago = int((datetime.now() - timedelta(hours=1)).timestamp())
    _console_send(client)
    out = _check(client, cloud, cloud_record(), cloud_record(
        id="865510083360422_00000001", timestamp=hour_ago, val1="AA55"))
    page = html.unescape(out["html"])
    assert page.count("sent elsewhere") == 1 and "AA55" in page
    assert page.index("AA55") < page.index("DBD76B82D531")      # oldest first


def test_a_record_of_another_type_is_not_taken_for_our_send(client, cloud):
    _console_send(client)
    out = _check(client, cloud, cloud_record(type=12))
    assert out["latest_stage"] == "accepted"
    assert "sent elsewhere" in html.unescape(out["html"])


def test_two_quick_sends_each_get_their_own_record(client, cloud):
    from ota_analytics import db

    _console_send(client, val1="AAAA")
    _console_send(client, val1="BBBB")
    now = int(datetime.now().timestamp())
    _check(client, cloud, cloud_record(id="R1", timestamp=now),
           cloud_record(id="R2", timestamp=now + 1))
    rows = db.connect().execute("SELECT COUNT(DISTINCT task_id) FROM cota_command "
                                "WHERE task_id IS NOT NULL").fetchone()[0]
    assert rows == 2


def test_an_imei_already_in_the_map_is_not_overwritten(client, cloud):
    from ota_analytics import db

    _map(db.connect(), imei="111111111111111")
    out = _check(client, cloud, cloud_record())
    assert out["label"] == "ID: 14906 | IMEI: 111111111111111"
    assert db.connect().execute("SELECT COUNT(*) FROM cota_device").fetchone()[0] == 1


# The live record once the device answered (2026-10-06), with the answer's FTP host and
# credentials replaced — same shape, same control bytes, nothing real.
ANSWER = ("(GET FTP_SETTINGS:203.0.113.6,1111,/mnt/vol1/sftp/x,user,secret,15,15:#j\x8e"
          "\u0016;Source:listener.example:4001)*EF")


def test_the_devices_answer_turns_the_ticks_green(client, cloud):
    from ota_analytics import db

    _console_send(client)
    out = _check(client, cloud, cloud_record())                    # status 0, response null
    assert out["latest_stage"] == "api"
    assert 'data-for="the device"' in out["html"]                  # still watching, for the device

    out = _check(client, cloud, cloud_record(status=1, response=ANSWER))
    assert out["latest_stage"] == "device"
    page = html.unescape(out["html"])
    assert 'class="tick tick-device"' in page and 'class="waiting"' not in page
    assert "Response via device" in page and "device answered" in page
    # Control bytes are shown, not rendered as garbage.
    assert "15:#j\\x8e\\x16;Source:listener.example:4001)*EF" in page
    row = db.connect().execute("SELECT status, device_response, device_response_at "
                               "FROM cota_command").fetchone()
    assert row["status"] == 1 and row["device_response"] == ANSWER and row["device_response_at"]


def test_a_devices_answer_is_kept_and_its_time_does_not_move(client, cloud):
    from ota_analytics import db

    _console_send(client)
    _check(client, cloud, cloud_record(status=1, response=ANSWER))
    conn = db.connect()
    first = conn.execute("SELECT device_response_at FROM cota_command").fetchone()[0]
    with conn:
        conn.execute("UPDATE cota_command SET device_response_at = '2026-10-06 11:29:00'")
    _check(client, cloud, cloud_record(status=1, response=None))   # an answer never disappears
    row = conn.execute("SELECT device_response, device_response_at FROM cota_command").fetchone()
    assert first and row[0] == ANSWER and row[1] == "2026-10-06 11:29:00"


def test_an_answer_with_a_readable_time_uses_it(client, cloud):
    from ota_analytics import db

    _console_send(client)
    _check(client, cloud, cloud_record(status=1, response="OK", responseTime=1791266400))
    at = db.connect().execute("SELECT device_response_at FROM cota_command").fetchone()[0]
    assert at == datetime.fromtimestamp(1791266400).strftime("%Y-%m-%d %H:%M:%S")


def test_the_live_send_reply_is_only_an_acknowledgement(client, cloud):
    """The send call answers {"msg": "Command send Successfully.."} — the one tick. The record
    comes from getGPRSCommand."""
    cloud.reply = '{"msg":"Command send Successfully.."}'
    body = client.get(_console_send(client).headers["location"]).text
    in_thread = _thread_of(body)
    assert 'class="tick tick-accepted"' in in_thread and "tick-api" not in in_thread
    assert "Command send Successfully.." in in_thread          # in the raw panel's template


def test_the_reply_window_covers_today_and_the_first_send(client, cloud):
    _console_send(client)
    client.post("/cota/console/check", data={"device": "14906"})
    _, start, end = cloud.checks[0]
    midnight = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    assert start <= int(midnight.timestamp()) and end > int(datetime.now().timestamp())


def test_a_body_with_no_records_is_shown_as_it_came(client, cloud):
    cloud.replies_body = '{"message": "no access"}'
    out = client.post("/cota/console/check", data={"device": "14906"}).json()
    assert "not with a list of command records" in out["html"]


def test_console_sends_build_one_job_per_device(client, cloud):
    from ota_analytics import db

    _console_send(client)
    _console_send(client, cmd_type="12", val1="60")
    rows = db.connect().execute("""
        SELECT j.name, j.source_file, t.seq, t.cmd_type FROM cota_task t
        JOIN cota_job j ON j.id = t.job_id ORDER BY t.seq
    """).fetchall()
    assert [tuple(r) for r in rows] == [("Console · device 14906", "console:14906", 1, 36),
                                        ("Console · device 14906", "console:14906", 2, 12)]
    # A console's sends live in its conversation on Configure, not among the jobs.
    assert "Console · device 14906" not in client.get("/cota").text


def test_a_console_send_needs_a_sign_in(client, cloud):
    from ota_analytics import cota

    cota.forget_token()
    reply = _console_send(client)
    assert reply.status_code == 200 and "Not signed in to the cloud" in reply.text
    assert 'value="DBD76B82D531"' in reply.text            # the draft is kept
    assert cloud.sent == []


def test_a_rejected_token_is_recorded_on_that_command_and_not_resent(client, cloud):
    from ota_analytics import cota, db

    cloud.error = cota.CotaError("The cloud rejected the token (HTTP 401). It has most likely "
                                 "expired.")
    reply = _console_send(client)
    assert reply.status_code == 303
    body = client.get(reply.headers["location"]).text
    assert "Not accepted" in body and "rejected the token" in body
    assert 'class="tick tick-failed"' in body
    assert [r[0] for r in db.connect().execute("SELECT state FROM cota_task")] == ["send_failed"]

    cloud.error = None
    _console_send(client, val1="NEXT")
    assert [b["val1"] for b in cloud.sent] == ["NEXT"]     # the failed one did not go out too


def test_the_device_card_says_what_web_fota_knows(client, cloud):
    from ota_analytics import db

    conn = db.connect()
    status = conn.execute("SELECT status FROM device WHERE imei = '444'").fetchone()[0]
    _map(conn, imei="444")
    body = client.get("/cota/console?device=14906").text
    assert "From Web FOTA" in body
    assert f"<strong>{status}</strong>" in body


def test_a_device_not_in_the_map_says_why_there_is_no_card(client, cloud):
    body = client.get("/cota/console?device=777").text
    assert "IMEI not known." in body and 'href="/cota/devices">Load the device map</a>' in body
    assert "ID: 777" in body


@pytest.mark.parametrize("text, reads", [
    ("SET TIMERS:KEEP,KEEP,FAIL,KEEP,KEEP,KEEP FAILED#123", "failed"),
    ("GET TIMERS:FOTA:6;HLTH:3600;IGNOFF:1800;SOS:300;TAMPER:60#123", None),
    ("SET BLEMAC:DBD76B82D531 OK", "ok"),
    ("ERROR: invalid parameter", "failed"),
])
def test_a_device_answer_reads_by_its_own_words(text, reads):
    from ota_analytics import cota

    assert cota.reading(text) == reads


def test_the_counts_match_the_portals_for_the_same_two_commands(client, cloud):
    """The portal's COTA screen for device 14906 on 2026-10-06: Total Count 2, Pending Count 1 —
    one answered (status 1) at 11:28:46, one pending (status 0) at 14:05:14."""
    first = int(datetime.now().timestamp()) - 3
    _console_send(client)
    _console_send(client, val1="SECOND")
    out = _check(client, cloud,
                 cloud_record(timestamp=first, status=1, response=ANSWER),
                 cloud_record(id="865510083360422_6AC4B2C2", timestamp=first + 2,
                              val1="245A454ED731323334D7DAD76F4BD76AC4B2C2D74145D1"))
    counts = re.sub(r"\s+", " ", html.unescape(out["counts_html"]))
    assert "Total <b>2</b>" in counts
    assert "Waiting for device <b>1</b>" in counts and "Answered <b>1</b>" in counts
    assert out["latest_stage"] == "api"              # the newest is still waiting


def test_an_unnamed_command_type_falls_back_to_its_number():
    from ota_analytics import cota

    assert cota.command_name(36) == "Zenithra Command"
    assert cota.command_name(12) == "Command 12"


# ─── Configure: the cleaned-up thread, the side panel, the range and the export ──

def test_the_legend_and_the_api_bubble_are_gone_and_the_raw_lives_in_the_side_panel(client, cloud):
    _console_send(client)
    _check(client, cloud, cloud_record())
    body = client.get("/cota/console?device=14906").text
    assert 'class="tick-legend"' not in body and "Held as" not in body
    assert '<aside class="console-side" id="console-side"' in body
    assert re.search(r'<template id="raw-t\d+">', body)              # one per command
    assert "Raw request — saveCOTAConfig" in body
    assert 'data-open-tab="device"' in body                          # the header status line


def test_the_header_says_whether_the_device_can_answer(client, cloud):
    from ota_analytics import db

    conn = db.connect()
    status = conn.execute("SELECT status FROM device WHERE imei = '444'").fetchone()[0]
    _map(conn, imei="444")
    head = client.get("/cota/console?device=14906").text.split('class="thread-who"', 1)[1]
    head = re.sub(r"\s+", " ", head.split("</button>", 1)[0])
    assert f"{status} · last seen" in head


def test_the_default_range_is_today_midnight_to_2359(client, cloud):
    body = client.get("/cota/console?device=14906").text
    today = datetime.now().strftime("%d-%m-%Y")
    assert f'name="from" class="mono dt-text" value="{today} 00:00"' in body
    assert f'name="to" class="mono dt-text" value="{today} 23:59"' in body


def test_the_range_filters_the_thread(client, cloud):
    from ota_analytics import db

    _console_send(client, val1="YESTERDAYS")
    conn = db.connect()
    yesterday = datetime.now() - timedelta(days=1)
    with conn:
        conn.execute("UPDATE cota_task SET sent_at = ?",
                     (yesterday.strftime("%Y-%m-%d 10:00:00"),))
    _console_send(client, val1="TODAYS")
    today_view = _thread_of(client.get("/cota/console?device=14906").text)
    assert "TODAYS" in today_view and "YESTERDAYS" not in today_view
    day = yesterday.strftime("%d-%m-%Y")
    old_view = _thread_of(client.get(
        f"/cota/console?device=14906&from={day} 00:00&to={day} 23:59").text)
    assert "YESTERDAYS" in old_view and "TODAYS" not in old_view


def test_a_range_over_15_days_is_cut_and_a_backwards_one_refused(client, cloud):
    body = client.get("/cota/console?device=14906&from=01-09-2026 00:00&to=30-09-2026 23:59").text
    assert "limited to 15 days" in body and 'value="15-09-2026 23:59"' in body
    body = client.get("/cota/console?device=14906&from=06-10-2026 10:00&to=05-10-2026 10:00").text
    assert "before its start" in body


def test_after_a_send_the_start_stays_and_the_end_moves_to_today(client, cloud):
    # Relative to today: a fixed date aged past MAX_RANGE_DAYS on 16-10-2026 and the start was
    # then rightly moved up — the case below is a start the range still holds.
    start = (datetime.now() - timedelta(days=3)).strftime("%d-%m-%Y") + " 06:00"
    reply = client.post("/cota/console/send", data={
        "device_id": "14906", "device_type": "124", "cmd_type": "36", "val1": "X1",
        "from": start}, follow_redirects=False)
    location = html.unescape(reply.headers["location"])
    from urllib.parse import parse_qs, urlsplit
    query = parse_qs(urlsplit(location).query)
    assert query["from"] == [start]
    assert query["to"] == [datetime.now().strftime("%d-%m-%Y") + " 23:59"]


def test_a_check_asks_the_cloud_for_the_pages_range(client, cloud):
    client.post("/cota/console/check", data={"device": "14906", "from": "05-10-2026 00:00",
                                             "to": "05-10-2026 23:59"})
    _, start, end = cloud.checks[0]
    assert start == int(datetime(2026, 10, 5).timestamp())
    assert end == int(datetime(2026, 10, 5, 23, 59, 59).timestamp())


def test_auto_check_is_three_tries_thirty_seconds_apart(client, cloud):
    from ota_analytics import cota

    assert (cota.CONSOLE_AUTO_EVERY, cota.CONSOLE_AUTO_TRIES) == (30, 3)
    body = client.get(_console_send(client).headers["location"]).text
    assert 'data-auto-every="30" data-auto-tries="3"' in body
    assert re.search(r'data-last-sent="\d{13}"', body)


def test_the_conversation_exports_as_csv_and_excel(client, cloud):
    _console_send(client)
    _check(client, cloud, cloud_record(status=1, response=ANSWER))
    csv_reply = client.get("/cota/console/export?device=14906&format=csv")
    assert csv_reply.status_code == 200
    assert "attachment" in csv_reply.headers["content-disposition"]
    text = csv_reply.content.decode("utf-8-sig")
    assert text.splitlines()[0] == ("IMEI,Device ID,Sent at,Value sent,Cloud val1,Cloud status,"
                                    "Response via API seen,Response via device,Answer seen")
    for dropped in ("Sent from", "Command", "Type", "Stage", "Send HTTP", "Send reply", "Cloud id"):
        assert f",{dropped}," not in f",{text.splitlines()[0]},"
    assert "15:#j\\x8e\\x16;Source" in text                 # control bytes as on screen

    xlsx_reply = client.get("/cota/console/export?device=14906&format=xlsx")
    from io import BytesIO
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(xlsx_reply.content))
    assert book.sheetnames == ["Commands", "Source"]
    source = {r[0]: r[1] for r in book["Source"].iter_rows(values_only=True)}
    assert source["Device"] == "ID: 14906 | IMEI: 865510083360422"
    assert source["Commands"] == "1" and source["Answered by the device"] == "1"


def test_native_controls_and_scrollbars_follow_the_theme():
    """A white scrollbar on the dark theme: the page never told the browser its colour scheme,
    so native controls drew themselves light. Every theme state now declares one."""
    from ota_analytics import config

    css = (config.resource("ota_analytics", "web") / "static" / "app.css").read_text(encoding="utf-8")
    assert css.count("color-scheme: dark;") >= 2 and css.count("color-scheme: light;") >= 2
    assert "scrollbar-color: var(--line) transparent" in css


def test_every_icon_a_template_asks_for_exists():
    """A missing icon renders as an empty outline with no error — the Export, Raw and panel
    buttons went blank that way. So every name used must be in nav.ICONS."""
    from ota_analytics import config, nav

    templates = config.resource("ota_analytics", "web") / "templates"
    used = set()
    for path in templates.glob("*.html"):
        text = path.read_text(encoding="utf-8")
        used |= set(re.findall(r"icon\('([\w-]+)'\)", text))
        used |= set(re.findall(r"nav_icons\.([\w]+)", text))
        used |= set(re.findall(r"nav_icons\['([\w-]+)'\]", text))
    used |= {item.icon for module in nav.MODULES for item in module.items}
    assert used and not (used - set(nav.ICONS)), f"no such icon: {sorted(used - set(nav.ICONS))}"


def test_the_console_actions_have_names(client, cloud):
    body = re.sub(r"\s+", " ", client.get("/cota/console?device=14906").text)
    for name in ("Refresh", "Export", "Raw", "Device", "Fold"):
        assert f"<span>{name}</span>" in body, f"{name} has no visible label"


def test_a_dev_copy_says_when_its_code_is_older_than_the_files(client, monkeypatch):
    from ota_analytics import api, config

    monkeypatch.setattr(config, "CHANNEL", "dev")
    monkeypatch.setattr(api, "_STARTED_WITH_CODE", 0.0)
    assert 'id="code-stale"' in client.get("/cota/console").text
    monkeypatch.setattr(api, "_STARTED_WITH_CODE", api._code_mtime())
    assert 'id="code-stale"' not in client.get("/cota/console").text
    monkeypatch.setattr(config, "CHANNEL", "release")
    monkeypatch.setattr(api, "_STARTED_WITH_CODE", 0.0)
    assert 'id="code-stale"' not in client.get("/").text        # a release never shows it


def test_no_bare_header_rule_can_make_other_headings_sticky():
    """`header { position: sticky }` applied to every <header>, and the Configure thread's
    heading painted over the top bar when the page scrolled. The top bar is styled by class."""
    from ota_analytics import config

    css = (config.resource("ota_analytics", "web") / "static" / "app.css").read_text(encoding="utf-8")
    assert not re.search(r"^header\s*[{,]", css, re.M)
    templates = config.resource("ota_analytics", "web") / "templates"
    users = [p.name for p in templates.glob("*.html") if "<header" in p.read_text(encoding="utf-8")]
    assert users == ["base.html"]


# ─── Configure: the period is the cloud's whole conversation for it ────────

def _check_range(client, cloud, day: datetime, *records, body=None):
    cloud.replies_body = body if body is not None else json.dumps({"data": list(records)})
    d = day.strftime("%d-%m-%Y")
    return client.post("/cota/console/check", data={"device": "14906", "from": f"{d} 00:00",
                                                    "to": f"{d} 23:59"}).json()


def test_a_past_day_shows_the_whole_conversation_the_cloud_holds(client, cloud):
    """Nothing was sent from this copy that day; the cloud still holds two commands sent from
    the portal, both answered. Loading the day shows both, with their answers."""
    yesterday = (datetime.now() - timedelta(days=1)).replace(hour=10, minute=0, second=0)
    t = int(yesterday.timestamp())
    out = _check_range(client, cloud, yesterday,
                       cloud_record(id="R_A", timestamp=t, status=1, response="GET A OK"),
                       cloud_record(id="R_B", timestamp=t + 600, status=1, response="GET B OK"))
    assert out["ok"] and out["message"].startswith("Checked ")
    page = html.unescape(out["html"])
    assert page.count("sent elsewhere") == 2
    assert "GET A OK" in page and "GET B OK" in page
    assert page.index("GET A OK") < page.index("GET B OK")


def test_opening_a_period_loads_it_unless_it_was_just_loaded(client, cloud):
    d = datetime.now().strftime("%d-%m-%Y")
    url = f"/cota/console?device=14906&from={d} 00:00&to={d} 23:59"
    body = client.get(url).text
    assert 'data-fresh="0"' in body and "Not checked yet" in body
    _check_range(client, cloud, datetime.now(), cloud_record())
    body = re.sub(r"\s+", " ", client.get(url).text)
    assert 'data-fresh="1"' in body
    assert re.search(r'id="check-status"[^>]*>Checked \d{2}:\d{2}:\d{2}<', body)
    assert "records in the cloud" not in body                    # when, not how many


def test_a_record_that_leaves_the_cloud_is_marked_and_comes_back_unmarked(client, cloud):
    from ota_analytics import db

    now = int(datetime.now().timestamp())
    keep, drop = cloud_record(id="KEEP", timestamp=now - 60), cloud_record(id="DROP", timestamp=now - 30)
    _check_range(client, cloud, datetime.now(), keep, drop)
    out = _check_range(client, cloud, datetime.now(), keep)                # deleted in the portal
    assert "1 no longer in the cloud" in out["message"]
    page = html.unescape(out["html"])
    assert "no longer in the cloud since" in page and 'gone"' in page
    assert "No longer in the cloud <b>1</b>" in re.sub(r"\s+", " ", out["counts_html"])
    row = db.connect().execute("SELECT missing_since FROM cota_command WHERE cloud_id = 'DROP'"
                               ).fetchone()
    assert row[0]

    _check_range(client, cloud, datetime.now(), keep, drop)                # it is back
    assert db.connect().execute("SELECT missing_since FROM cota_command WHERE cloud_id = 'DROP'"
                                ).fetchone()[0] is None


def test_an_empty_day_marks_what_it_no_longer_holds(client, cloud):
    from ota_analytics import db

    _check_range(client, cloud, datetime.now(), cloud_record())
    out = _check_range(client, cloud, datetime.now(), body='{"data": []}')
    assert out["ok"] and "1 no longer in the cloud" in out["message"]
    assert db.connect().execute("SELECT missing_since IS NOT NULL FROM cota_command").fetchone()[0]


def test_only_the_asked_window_and_a_readable_answer_can_mark_anything(client, cloud):
    from ota_analytics import db

    yesterday = datetime.now() - timedelta(days=1)
    _check_range(client, cloud, yesterday, cloud_record(id="OLD", timestamp=int(yesterday.timestamp())))
    _check_range(client, cloud, datetime.now(), cloud_record(id="NEW"))   # another day's window
    _check_range(client, cloud, datetime.now(), body='{"message": "busy"}')  # says nothing
    rows = dict(db.connect().execute("SELECT cloud_id, missing_since FROM cota_command").fetchall())
    assert rows == {"OLD": None, "NEW": None}


def test_schema_v12_records_when_a_command_left_the_cloud(client):
    # `client` points the database at a temp file — without it this would open the real one.
    from ota_analytics import db

    conn = db.connect()
    assert db.SCHEMA_VERSION >= 12
    assert "missing_since" in {r["name"] for r in conn.execute("PRAGMA table_info(cota_command)")}


def test_model_and_type_from_the_edit_panel_go_out_with_the_send(client, cloud):
    """They left the composer, but they are still the composer's fields (form="composer"), so a
    send carries them — and a changed type survives the reload after it."""
    reply = client.post("/cota/console/send", data={"device_id": "14906", "device_type": "125",
                                                    "cmd_type": "12", "val1": "60"},
                        follow_redirects=False)
    assert cloud.sent == [{"deviceType": [125], "deviceList": [14906], "type": 12, "val1": "60"}]
    assert "cmd=12" in reply.headers["location"]
    head = client.get(reply.headers["location"]).text.split('class="thread-title-row"', 1)[1]
    assert 'value="12"' in head.split("</details>", 1)[0]
    assert "model 125 · type 12" in head                        # the pill: not the defaults


def test_the_console_fills_the_window_and_the_panes_share_a_divider(client, cloud):
    body = client.get("/cota/console?device=14906").text
    assert '<div class="page fit-window">' in body
    from ota_analytics import config
    css = (config.resource("ota_analytics", "web") / "static" / "app.css").read_text(encoding="utf-8")
    assert ".page.fit-window { height: 100vh; }" in css
    assert ".console-list { border-radius: 10px 0 0 10px;" in css


# ─── an expired token shows on the cloud chip ──────────────────────────────

@pytest.fixture
def real_client_cloud(cota_env, monkeypatch):
    """The real cota.Client, over a fake network: `answer["status"]` decides what the cloud
    says, so the chip is driven by exactly the code that runs against the live cloud."""
    import httpx

    from ota_analytics import cota

    answer = {"status": 401, "body": '{"message": "Unauthorized"}'}

    def handler(request):
        return httpx.Response(answer["status"], text=answer["body"])

    real = httpx.Client
    monkeypatch.setattr(httpx, "Client",
                        lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    monkeypatch.setattr(cota, "CALL_INTERVAL_SECONDS", 0)
    cota.save_token("an-old-token-that-has-expired")
    return answer


def _chip(body: str) -> str:
    chip = body.split('id="cota-signin-chip"', 1)[0].rsplit("<a ", 1)[1] + \
        body.split('id="cota-signin-chip"', 1)[1].split("</a>", 1)[0]
    return re.sub(r"\s+", " ", chip)


def test_a_rejected_token_turns_the_cloud_chip_red_at_once(client, real_client_cloud):
    assert "Cloud: signed in" in _chip(client.get("/cota/console?device=14906").text)

    out = client.post("/cota/console/check", data={"device": "14906"}).json()
    assert not out["ok"] and "rejected the token (HTTP 401)" in out["message"]
    assert "cota token set" not in out["message"]
    assert out["cloud"]["level"] == "error" and out["cloud"]["label"] == "Cloud: session expired"

    body = client.get("/cota/console?device=14906").text
    chip = _chip(body)
    assert "status-error" in chip and "Cloud: session expired" in chip and "HTTP 401" in chip
    assert "Sign in again" in body                               # the composer says it too
    assert "rejected the saved token" in client.get("/cota/signin").text


def test_a_rejected_send_also_shows_on_the_chip(client, real_client_cloud):
    reply = client.post("/cota/console/send", data={"device_id": "14906", "device_type": "124",
                                                    "cmd_type": "36", "val1": "X"},
                        follow_redirects=False)
    body = client.get(reply.headers["location"]).text
    assert "Cloud: session expired" in _chip(body) and "Not accepted" in body


def test_a_fresh_token_or_an_accepted_call_clears_it(client, real_client_cloud):
    client.post("/cota/console/check", data={"device": "14906"})
    client.post("/cota/signin", data={"cloud": "intouch", "method": "token",
                                      "token": "a-fresh-token-0123456789"})
    assert "Cloud: signed in" in _chip(client.get("/cota").text)

    client.post("/cota/console/check", data={"device": "14906"})           # rejected again
    assert "Cloud: session expired" in _chip(client.get("/cota").text)
    real_client_cloud.update(status=200, body='{"data": []}')              # the cloud relents
    out = client.post("/cota/console/check", data={"device": "14906"}).json()
    assert out["ok"] and out["cloud"]["label"] == "Cloud: signed in"


def test_the_sign_in_page_links_the_portal_but_never_posts_a_password_to_it(client, cota_env):
    """The CTVMS portal's sign-in page is where a fresh token comes from. It is not a login API:
    the part after '#' never reaches the server, so it must not be used as the login URL."""
    from ota_analytics import cota_connection

    body = client.get("/cota/signin").text
    assert 'href="https://ctvms.mappls.com/adminnextgen/#/login"' in body
    assert cota_connection.load().login_url == ""
    assert 'name="login_url" value=""' in body


# ─── Configure: sequences from the page ────────────────────────────────────

USER_SEQUENCE = "1. DAD76F4B\n2. DAD76C0A\n3. DDD76D66\n4. DBD76C0AD531D931D9322E35D9332E35D9332E35D931D93130D933"


class _LiveThread:
    """Stands in for a runner thread: alive, so restart recovery leaves the run alone — as it
    does in the app, where the thread is real."""

    def is_alive(self):
        return True


@pytest.fixture
def no_run_threads(monkeypatch):
    from ota_analytics import cota_run

    started, threads = [], {}

    def start(run_id):
        started.append(run_id)
        threads[run_id] = _LiveThread()

    monkeypatch.setattr(cota_run, "start", start)
    monkeypatch.setattr(cota_run, "_threads", threads)
    return started


def _start_sequence(client, text=USER_SEQUENCE):
    return client.post("/cota/console/run", data={"device_id": "14906", "device_type": "124",
                                                  "cmd_type": "36", "commands": text},
                       follow_redirects=False)


def test_a_sequence_starts_in_the_background_and_the_page_shows_it(client, cloud, no_run_threads):
    from ota_analytics import db

    reply = _start_sequence(client)
    assert reply.status_code == 303 and reply.headers["location"].endswith("#run")
    assert no_run_threads == [1]                                   # started, not run inline
    steps = db.connect().execute("SELECT val1, kind, state FROM cota_run_step ORDER BY seq").fetchall()
    assert [(s["kind"], s["state"]) for s in steps] == [("get", "queued"), ("get", "queued"),
                                                         ("clear", "queued"), ("set", "queued")]
    assert "3 · CLR SOS" in client.get(reply.headers["location"]).text
    body = client.get(reply.headers["location"]).text
    assert 'id="run-panel"' in body and "running · 0 of 4 done" in re.sub(r"\s+", " ", body)
    assert 'data-run-live="1"' in body
    assert cloud.sent == []                                         # nothing sent by the page


def test_a_bad_line_is_refused_and_nothing_starts(client, cloud, no_run_threads):
    body = _start_sequence(client, "DAD76F4B\nhello").text
    assert "line 2" in body and "not hexadecimal" in body
    assert no_run_threads == [] and 'value="' in body and "hello</textarea>" in body


def test_while_a_sequence_is_live_nothing_else_goes_to_the_device(client, cloud, no_run_threads):
    _start_sequence(client)
    hand = client.post("/cota/console/send", data={"device_id": "14906", "device_type": "124",
                                                   "cmd_type": "36", "val1": "X"})
    assert "A sequence is running for this device" in hand.text and cloud.sent == []
    second = _start_sequence(client, "DAD76F4B")
    assert "already running" in second.text and len(no_run_threads) == 1


def test_pause_resume_and_cancel_from_the_page(client, cloud, no_run_threads):
    from ota_analytics import db

    _start_sequence(client)
    conn = db.connect()
    client.post("/cota/console/run/control", data={"run_id": 1, "action": "pause"})
    assert conn.execute("SELECT control FROM cota_run").fetchone()[0] == "pause"
    with conn:                                       # as the runner does when it sees the request
        conn.execute("UPDATE cota_run SET state = 'paused', control = NULL")
    client.post("/cota/console/run/control", data={"run_id": 1, "action": "resume"})
    assert conn.execute("SELECT state FROM cota_run").fetchone()[0] == "running"
    assert no_run_threads == [1, 1]                  # resuming starts the runner again
    with conn:
        conn.execute("UPDATE cota_run SET state = 'paused'")
    client.post("/cota/console/run/control", data={"run_id": 1, "action": "cancel"})
    assert conn.execute("SELECT state FROM cota_run").fetchone()[0] == "cancelled"
    assert {r[0] for r in conn.execute("SELECT state FROM cota_run_step")} == {"cancelled"}


def test_the_run_status_reads_only_the_local_record(client, cloud, no_run_threads):
    _start_sequence(client)
    out = client.get("/cota/console/run-status?device=14906").json()
    assert out["ok"] and out["state"] == "running" and 'id="run-panel"' in out["run_html"]
    assert cloud.checks == [] and cloud.sent == []                # polled for free


def test_the_control_form_only_ever_returns_to_the_console(client, cloud, no_run_threads):
    _start_sequence(client)
    reply = client.post("/cota/console/run/control", data={"run_id": 1, "action": "pause",
                                                           "back": "https://evil.example/"},
                        follow_redirects=False)
    assert reply.headers["location"] == "/cota/console"



def test_a_sequence_whose_runner_is_gone_comes_back_paused_on_the_page(client, cloud, monkeypatch):
    """No thread behind a 'running' run means the app restarted mid-run: the page says so."""
    from ota_analytics import cota_run

    monkeypatch.setattr(cota_run, "start", lambda run_id: None)        # started, then lost
    monkeypatch.setattr(cota_run, "_threads", {})
    reply = _start_sequence(client)
    body = re.sub(r"\s+", " ", client.get(reply.headers["location"]).text)
    assert "paused · 0 of 4 done" in body and "The app was restarted" in body
    assert "Resume</button>" in body


# ─── Intouch COTA: jobs and groups ──────────────────────────────────────────

@pytest.fixture
def no_job_threads(monkeypatch):
    """A job's scheduler never runs here: the page only has to start it and draw its record."""
    from ota_analytics import cota_campaign

    started, threads = [], {}

    def start(campaign_id):
        started.append(campaign_id)
        threads[campaign_id] = _LiveThread()

    monkeypatch.setattr(cota_campaign, "start", start)
    monkeypatch.setattr(cota_campaign, "_threads", threads)
    return started


def _csv_upload(text: str, name: str = "devices.csv"):
    return {"file": (name, text.encode("utf-8"), "text/csv")}


def _job_preview(client, **form):
    data = {"name": "FTP check", "commands": "DAD76F4B", **form}
    return client.post("/cota/jobs/preview", data=data).text


def test_a_wait_after_a_command_goes_from_the_library_through_the_plan_to_the_job(client, cloud,
                                                                                no_job_threads):
    """10-10-2026: SET Force Sync Audio takes the device off the air; the next command must wait.
    Saved on the command, shown and changeable per line, counted in the plan, kept by the job."""
    reply = client.post("/cota/commands/saved", data={"name": "Force Sync Audio", "val1": "DBD76F85D531",
                                                      "wait_after": "180"}, follow_redirects=False)
    assert reply.status_code == 303
    library = re.sub(r"\s+", " ", client.get("/cota/commands").text)
    assert "3 min</span>" in library and 'data-wait="180"' in library
    sync, ftp = "DBD76F85D531", "DAD76F4B"
    # The form sends one wait per line; without them the saved command's is the default.
    default = re.sub(r"\s+", " ", _job_preview(client, devices="786", commands=f"{sync}\n{ftp}"))
    assert 'name="waits" value="180,0"' in default and "then wait 3 min" in default
    changed = re.sub(r"\s+", " ", client.post("/cota/jobs/preview", data={
        "name": "Sync", "devices": "786", "commands": f"{sync}\n{ftp}", "waits": ["300", "0"]}).text)
    assert 'name="waits" value="300,0"' in changed and 'data-waits="300,0"' in changed
    client.post("/cota/jobs/start", data={"device_ids": "786", "commands": f"{sync}\n{ftp}", "waits": "300,0"})
    assert db.connect().execute("SELECT waits FROM cota_campaign").fetchone()[0] == "[300, 0]"
    page = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert "5 min</span>" in page                                     # on the command summary
    again = client.get("/cota/jobs/new?from=1").text
    assert 'data-waits="300,0"' in again                               # Duplicate keeps them


def test_the_upload_template_is_the_clouds_own_device_list(client, cota_env):
    reply = client.get("/cota/groups/template.csv")
    assert reply.status_code == 200
    assert "cota_devices_template.csv" in reply.headers["content-disposition"]
    lines = reply.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "id,trackingCode"
    # The template shows trackingCode is optional by leaving one blank.
    assert any(line.endswith(",") for line in lines[1:])


def test_both_upload_forms_offer_the_template(client, cota_env):
    assert 'href="/cota/groups/template.csv"' in client.get("/cota/jobs/new").text
    body = client.get("/cota/devices").text           # Add device: the same list, with tags
    assert 'href="/cota/devices/template.csv"' in body and "Download template" in body


def test_a_group_name_never_reaches_a_script_string(client, cota_env):
    """HTML-escaping does not protect a value inside an onsubmit: the attribute is decoded
    before the script runs. The name goes through a data attribute instead."""
    name = "x');alert(1);('"
    body = client.post("/cota/groups", data={"name": name, "devices": "786"}).text
    # No inline script at all: the confirmation reads the name from data-name and sets it as text.
    assert "onsubmit=" not in body and "x');alert" not in body          # quotes always escaped
    assert 'data-name="x&#39;);alert(1);(&#39;"' in body
    assert 'data-confirm="Delete the group “{name}”?' in body


def test_a_preview_plans_the_job_and_sends_nothing(client, cloud, no_job_threads):
    body = re.sub(r"\s+", " ", _job_preview(client, devices="14906, 786",
                                            commands="DAD76F4B\nDAD76C0A"))
    assert "Plan" in body and "GET FTP_SETTINGS" in body
    assert "Start job — 2 devices" in body
    assert 'name="device_ids" value="14906,786"' in body
    assert cloud.sent == [] and cloud.checks == [] and no_job_threads == []


def test_the_commands_reach_start_with_their_line_breaks(client, cloud, no_job_threads):
    body = _job_preview(client, devices="786", commands="DAD76F4B\nDAD76C0A")
    assert re.search(r'<textarea name="commands" hidden>DAD76F4B\r?\nDAD76C0A</textarea>', body)


def test_a_preview_from_a_group_uses_the_groups_devices(client, cloud, no_job_threads):
    client.post("/cota/groups", data={"name": "Desk", "devices": "14906, 786"})
    gid = db.connect().execute("SELECT id FROM cota_group").fetchone()[0]
    body = _job_preview(client, group_id=str(gid))
    assert 'name="device_ids" value="14906,786"' in body
    assert f'<option value="{gid}" selected' in body


def test_a_preview_from_an_uploaded_csv(client, cloud, no_job_threads):
    body = client.post("/cota/jobs/preview", data={"commands": "DAD76F4B"},
                       files=_csv_upload("id,trackingCode\n786,\n14906,865510083360422\n")).text
    assert 'name="device_ids" value="786,14906"' in body          # the file's order is kept


def test_a_bad_command_blocks_start(client, cloud, no_job_threads):
    body = _job_preview(client, devices="786", commands="DAD76F4B\nhello")
    assert "Fix these commands first" in body and "line 2" in body
    assert "Start job" not in body


def test_without_a_sign_in_the_plan_says_so_instead_of_offering_start(client, cota_env):
    body = _job_preview(client, devices="786")
    assert "Not signed in to the cloud" in body and "Start job" not in body


def test_starting_a_job_runs_it_in_the_background(client, cloud, no_job_threads):
    reply = client.post("/cota/jobs/start", data={"name": "FTP check", "device_ids": "14906,786",
                                                  "commands": "DAD76F4B"}, follow_redirects=False)
    assert reply.status_code == 303 and reply.headers["location"] == "/cota/jobs/1"
    assert no_job_threads == [1] and cloud.sent == []              # started, not run inline
    body = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert "#1 · FTP check" in body and '<span data-tw="finished">0</span> of 2 device commands finished' in body
    assert "GET FTP_SETTINGS" in body and "<span>Pause</span></button>" in body
    assert "/cota/jobs/1/export?format=csv" in body
    assert 'data-unfold="786"' in body                               # each device's conversation
    listing = client.get("/cota").text
    assert 'href="/cota/jobs/1"' in listing and 'data-tip="One job runs at a time."' in listing


def test_a_job_page_redraws_from_the_local_record_only(client, cloud, no_job_threads):
    client.post("/cota/jobs/start", data={"device_ids": "786", "commands": "DAD76F4B"})
    out = client.get("/cota/jobs/1/status").json()
    assert out["ok"] and out["state"] == "running" and "device commands finished" in out["progress"]
    assert "Command summary" in out["html"] and "Command summary" not in out["progress"]
    assert cloud.checks == [] and cloud.sent == []
    assert client.get("/cota/jobs/99/status").json() == {"ok": False}
    assert client.get("/cota/jobs/99", follow_redirects=False).headers["location"] == "/cota"


def test_devices_in_progress_go_from_the_form_to_the_job_and_a_batch_is_held_to_200(
        client, cloud, no_job_threads):
    ids = ", ".join(str(n) for n in range(1, 451))
    body = re.sub(r"\s+", " ", _job_preview(client, devices=ids, in_progress="150", batch_size="999"))
    assert "<em>At a time</em><b>150</b> <small>the rest in line — 3 rounds</small>" in body
    assert 'name="in_progress" value="150"' in body and 'name="batch_size" value="200"' in body
    reply = client.post("/cota/jobs/start", data={"device_ids": ids, "commands": "DAD76F4B",
                                                  "in_progress": "150", "batch_size": "200"},
                        follow_redirects=False)
    assert reply.status_code == 303
    conn = db.connect()
    assert tuple(conn.execute("SELECT in_progress, batch_size FROM cota_campaign").fetchone()) == (150, 200)
    page = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert "150 at a time" in page and "300 devices in line" in page and ">In line</a>" in page
    # Each tile's share leads, its count beside it (the user, 08-10-2026): nothing sent yet.
    assert '<span data-tw="not_sent-pct">100</span>%<span class="tile-count" data-tw="not_sent-n">450</span>' in page
    assert '<span data-tw="failed-pct">0</span>%<span class="tile-count" data-tw="failed-n">0</span>' in page
    assert "on the first attempt</div>" not in page              # the tile's sub-line is gone
    # The title, actions and progress are the sticky top; the devices scroll under it.
    assert page.index('id="job-top"') < page.index('id="job-progress"') < page.index('id="job-live"')


def test_a_device_that_answered_nothing_is_marked_not_reachable(client, cloud, no_job_threads):
    client.post("/cota/jobs/start", data={"device_ids": "786, 14906", "commands": "DAD76F4B"})
    conn = db.connect()
    with conn:                                     # 786 answered; 14906 gave up unanswered
        conn.execute("UPDATE cota_campaign_result SET state = 'done', outcome = 'answered' "
                     "WHERE device_id = 786")
        conn.execute("UPDATE cota_campaign_result SET state = 'failed', outcome = 'not_delivered' "
                     "WHERE device_id = 14906")
    body = client.get("/cota/jobs/1?state=unreachable").text
    assert 'class="on">Not reachable</a>' in body
    assert body.count("pill-dev-unreachable") == 1 and 'data-unfold="14906"' in body
    assert 'data-unfold="786"' not in body                        # filtered out: it answered


def test_a_device_in_a_live_job_is_refused_by_the_next_plan(client, cloud, no_job_threads):
    client.post("/cota/jobs/start", data={"device_ids": "786", "commands": "DAD76F4B"})
    body = _job_preview(client, devices="786, 14906")
    assert "in a live" in body and "Start job" not in body


def test_pause_and_cancel_from_the_job_page(client, cloud, no_job_threads):
    client.post("/cota/jobs/start", data={"device_ids": "786", "commands": "DAD76F4B"})
    conn = db.connect()
    reply = client.post("/cota/jobs/1/control", data={"action": "pause"}, follow_redirects=False)
    assert reply.headers["location"] == "/cota/jobs/1"
    assert conn.execute("SELECT control FROM cota_campaign").fetchone()[0] == "pause"
    with conn:                                     # as the scheduler does when it sees it
        conn.execute("UPDATE cota_campaign SET state = 'paused', control = NULL")
    client.post("/cota/jobs/1/control", data={"action": "cancel"})
    assert conn.execute("SELECT state FROM cota_campaign").fetchone()[0] == "cancelled"
    assert "<span>Pause</span></button>" not in client.get("/cota/jobs/1").text


def test_the_job_export_has_a_row_per_device_and_command(client, cloud, no_job_threads):
    client.post("/cota/jobs/start", data={"device_ids": "14906,786",
                                          "commands": "DAD76F4B\nDAD76C0A"})
    reply = client.get("/cota/jobs/1/export?format=csv")
    assert reply.status_code == 200
    lines = reply.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Device ID,IMEI,Device type,FW version,Step,Command")   # firmware: 10-10-2026
    assert len(lines) == 1 + 2 * 2
    xlsx = client.get("/cota/jobs/1/export?format=xlsx")
    assert xlsx.status_code == 200 and xlsx.content[:4] == b"PK\x03\x04"


# ─── 2.0.1: the jobs list, a new job, the job page, the command library ─────

def _start_job(client, **over):
    data = {"name": "FTP check", "device_ids": "14906,786", "commands": "DAD76F4B\nDAD76C0A", **over}
    return client.post("/cota/jobs/start", data=data, follow_redirects=False)


def test_a_job_shows_its_clock_and_start_says_how_long_it_can_take(client, cloud, no_job_threads):
    """09-10-2026: start time, max time and remaining on the job; the max before Start."""
    preview = re.sub(r"\s+", " ", _job_preview(client, devices="786, 14906", commands="DAD76F4B\nDAD76C0A"))
    assert 'data-confirm-tone="go"' in preview and 'data-confirm-action="Start job"' in preview
    # 2 commands × 100 s and a tenth: 4 min — the job's hard limit (10-10-2026).
    # The confirmation is a formal summary: a sentence, then the figures as a table (10-10-2026).
    assert 'data-confirm="2 commands will be sent to 2 devices, in the order listed."' in preview
    summary = preview.split('<template id="start-summary">')[1].split("</template>")[0]
    assert "<dt>Maximum duration</dt><dd>4 min</dd>" in summary and "<dt>Stops by</dt>" in summary
    assert "<dt>Devices at a time</dt><dd>2</dd>" in summary and "<dt>Estimated duration</dt>" in summary
    assert "<em>May take up to</em> <b>4 min</b> <small>hard stop · ends by " in preview
    # The arithmetic is in the figure's hold-tip, not printed (10-10-2026).
    assert ('data-tip="1 round × 2 commands × 1 min 40 s (30 s for an answer, up to 3 attempts) '
            '= 3 min 20 s, + 10% → 4 min.') in preview
    assert 'class="plan-max-sum"' not in preview and "Each command</dt>" not in preview
    client.post("/cota/jobs/start", data={"device_ids": "786, 14906", "commands": "DAD76F4B\nDAD76C0A"})
    assert db.connect().execute("SELECT validity_hours * 60 FROM cota_campaign").fetchone()[0] == 4
    page = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert "<em>Started</em>" in page and "<em>Max time</em><b>4 min</b>" in page
    assert "<em>Remaining</em><b>up to 3 min 20 s</b>" in page


def test_a_job_its_time_limit_stopped_says_so_and_offers_the_rest(client, cloud, no_job_threads):
    client.post("/cota/jobs/start", data={"device_ids": "1, 2, 3", "commands": "DAD76F4B"})
    conn = db.connect()
    with conn:                                     # the limit came: 1 finished, 1 mid-way, 1 never began
        conn.execute("UPDATE cota_campaign SET state = 'done'")
        conn.execute("UPDATE cota_campaign_device SET state = 'done' WHERE device_id = 1")
        conn.execute("UPDATE cota_campaign_device SET state = 'expired' WHERE device_id = 2")
        conn.execute("UPDATE cota_campaign_device SET state = 'not_started' WHERE device_id = 3")
    page = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert ">time limit</span>" in page and ">done</span>" not in page
    assert "Stopped at its 2-minute time limit — 1 mid-sequence and 1 not started." in page
    assert 'href="/cota/jobs/new?from=1&which=rest"' in page and "Continue the rest — 2" in page
    form = re.sub(r"\s+", " ", client.get("/cota/jobs/new?from=1&which=rest").text)
    assert "the devices its time limit stopped (2)" in form
    assert ">time limit</span>" in client.get("/cota").text


def test_the_first_visit_of_a_new_day_never_shows_yesterdays_jobs(client, cota_env, monkeypatch):
    """10-10-2026: the Overview computed its figures before the day's clearing ran, so the first
    visit of a new day showed yesterday's jobs as today's, and the second showed nothing."""
    from ota_analytics import cota_campaign
    conn = db.connect()
    cid = cota_campaign.create(conn, name="yesterday", device_ids=[5001], commands_text="DAD76F4B")
    with conn:
        conn.execute("UPDATE cota_campaign_result SET attempts = 1, state = 'done', outcome = 'answered'")
        conn.execute("UPDATE cota_campaign SET created_at = '2026-01-01 10:00:00', state = 'done' WHERE id = ?", (cid,))
    monkeypatch.setattr(cota_campaign, "_purged_this_process", False)   # a fresh process, a new day
    assert "No commands sent today." in client.get("/cota/overview").text
    assert 'href="/cota/jobs/1"' not in client.get("/cota").text


def test_the_cota_overview_reads_the_day_and_asks_web_fota_about_the_silent(client, cota_env):
    assert "No commands sent today." in client.get("/cota/overview").text
    from ota_analytics import cota, cota_campaign
    conn = db.connect()
    cota_campaign.create(conn, name="t", device_ids=[5001, 5002], commands_text="DAD76F4B")
    now = cota._now()
    with conn:
        conn.execute("UPDATE cota_campaign_result SET attempts = 1, state = 'done', outcome = 'answered', "
                     "answered_attempt = 1, answer_seconds = 18, finished_at = ? WHERE device_id = 5001", (now,))
        conn.execute("UPDATE cota_campaign_result SET attempts = 3, state = 'failed', "
                     "outcome = 'not_delivered', finished_at = ? WHERE device_id = 5002", (now,))
        # 5002's IMEI is a Web FOTA device of this fixture, last seen two hours ago.
        cota.put_device(conn, 5002, imei="222", source="test")
    body = re.sub(r"\s+", " ", client.get("/cota/overview").text)
    for heading in ("Answer rate by hour", "Today at a glance", "Needs attention", "Commands today",
                    "By firmware"):
        assert heading in body
    assert "Held by the cloud" in body
    # Taken off at the user's request (10-10-2026).
    assert "Never answered — last ping" not in body and "Why commands failed" not in body
    assert 'href="/cota/overview"' in body                     # first in Intouch COTA's rail


def test_jobs_opens_on_the_list_and_new_job_is_its_own_page(client, cota_env):
    body = client.get("/cota").text
    assert 'href="/cota/jobs/new"' in body and "New job</span>" in body
    assert 'action="/cota/jobs/preview"' not in body                 # the form is not on the list
    assert "No jobs today" in body
    new = client.get("/cota/jobs/new").text
    assert 'action="/cota/jobs/preview"' in new and 'href="/cota"' in new and "All jobs" in new


def test_the_list_shows_each_jobs_progress_and_sequence(client, cloud, no_job_threads):
    _start_job(client)
    body = re.sub(r"\s+", " ", client.get("/cota").text)
    assert 'class="mini-progress"' in body and "0.0%" in body
    assert "<span>GET FTP_SETTINGS</span>" in body and "<span>GET 6C0A</span>" in body


def test_a_job_page_goes_back_to_the_list_and_titles_its_progress(client, cloud, no_job_threads):
    _start_job(client)
    body = client.get("/cota/jobs/1").text
    assert 'class="back-link" href="/cota"' in body
    assert re.search(r"<title>(DEV · )?\(0\.0%\) #1 FTP check — Jobs</title>", body)


def test_the_device_table_uses_the_users_headers(client, cloud, no_job_threads):
    _start_job(client)
    body = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    heads = re.findall(r"<th[^>]*>(.*?)</th>", body.split('class="job-devices data-grid"')[1].split("</thead>")[0])
    labels = [re.sub(r"<[^>]+>", "", h).strip() for h in heads]
    assert labels == ["Conversation", "S.No.", "Device ID", "IMEI", "State", "Command",
                      "Now at/Total", "Attempt", "Next", "Answered", "Failed", "Time to answer",
                      "Last answer"]
    row = body.split('class="job-devices data-grid"')[1].split("<tbody>")[1].split("</tr>")[0]
    assert ">1</td>" in row and ">GET FTP_SETTINGS</td>" in row and ">1/2</td>" in row


def test_the_job_device_table_sorts_on_its_columns(client, cloud, no_job_threads):
    """10-10-2026: every header sorts but Next (worked out per row, live); the order survives the
    live redraw, the filters and the pager."""
    _start_job(client, device_ids="14906,786,5")
    ids = lambda html: [int(x) for x in re.findall(r'data-unfold="(\d+)"', html)]
    assert ids(client.get("/cota/jobs/1").text) == [14906, 786, 5]                  # job order
    page = client.get("/cota/jobs/1?sort=id&dir=desc").text
    assert ids(page) == [14906, 786, 5][::1] and ids(client.get("/cota/jobs/1?sort=id").text) == [5, 786, 14906]
    assert 'aria-sort="ascending"' in client.get("/cota/jobs/1?sort=id").text
    assert 'data-query="state=&amp;q=&amp;sort=id&amp;dir=desc&page=1"' in page
    live = client.get("/cota/jobs/1/status?sort=id&dir=asc").json()["html"]
    assert ids(live) == [5, 786, 14906]
    assert 'href="/cota/jobs/1?state=waiting&amp;q=&amp;sort=id&amp;dir=desc"' in page   # filters keep it


def test_a_device_unfolds_to_its_conversation_in_the_job(client, cloud, no_job_threads):
    _start_job(client)
    closed = client.get("/cota/jobs/1/status").json()["html"]
    assert 'data-unfold="786"' in closed and "Open in Configure" not in closed
    opened = client.get("/cota/jobs/1/status?open=786").json()
    assert "Open in Configure" in opened["html"] and "Not sent yet</span>" in opened["html"]
    assert opened["percent"] == 0.0
    assert cloud.checks == [] and cloud.sent == []                    # redrawn from the record


def test_the_answer_wait_and_time_limit_are_job_settings(client, cloud, no_job_threads):
    _start_job(client, answer_wait_seconds="45", time_limit_minutes="20")
    row = db.connect().execute("SELECT answer_wait_seconds, validity_hours FROM cota_campaign").fetchone()
    assert row[0] == 45 and abs(row[1] - 20 / 60) < 1e-9
    body = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    # The settings are in the title's tooltip, not on the page (the user, 08-10-2026).
    assert re.search(r'<h2 data-tip="[^"]*answer wait 45 s · time limit 20 min, ends by', body)
    assert 'class="job-meta"' not in body and 'class="step-chain"' not in body
    preview = client.post("/cota/jobs/preview", data={"devices": "786", "commands": "DAD76F4B",
                                                      "answer_wait_seconds": "40"}).text
    assert re.search(r'name="answer_wait_seconds"[^>]*value="40"', preview)
    assert "(40 s for an answer, up to 3 attempts)" in preview and "hard stop · ends by" in preview
    # The fold is titled alone; a setting moved off its default is marked (10-10-2026).
    assert ">Advanced settings</span> <span class=\"pill pill-changed\">Changed</span>" in preview
    assert "pill-changed" not in client.get("/cota/jobs/new").text


def test_a_plan_that_cannot_fit_its_time_limit_says_so(client, cloud, no_job_threads):
    ids = ", ".join(str(n) for n in range(100000, 101500))
    body = re.sub(r"\s+", " ", client.post("/cota/jobs/preview", data={
        "devices": ids, "commands": "DAD76F4B\nDAD76C0A\nDDD76D66", "time_limit_minutes": "5"}).text)
    assert "Your 5-minute limit is shorter than this job's max time" in body and "Start job" in body


def test_the_preview_says_how_long_it_may_take_and_the_sum_behind_it(client, cloud, no_job_threads):
    """The user's example (10-10-2026): 2,000 devices × 6 commands, 200 at a time — 10 rounds ×
    6 × 100 s = 100 min, and a tenth: 110 min, the job's hard stop instead of a fixed hour."""
    ids = ", ".join(str(n) for n in range(100000, 102000))
    body = re.sub(r"\s+", " ", client.post("/cota/jobs/preview", data={
        "devices": ids, "commands": "\n".join(["DAD76F4B"] * 6)}).text)
    assert "<em>May take up to</em> <b>1 h 50 min</b> <small>hard stop · ends by " in body
    assert ("10 rounds × 6 commands × 1 min 40 s (30 s for an answer, up to 3 attempts) = 1 h 40 min, "
            "+ 10% → 1 h 50 min.") in body
    typed = re.sub(r"\s+", " ", client.post("/cota/jobs/preview", data={
        "devices": ids, "commands": "\n".join(["DAD76F4B"] * 6), "time_limit_minutes": "60"}).text)
    assert "<em>Stops at</em> <b>1 h</b> <small>your limit · its max time 1 h 50 min" in typed


def test_a_job_that_cannot_finish_today_is_held_with_the_number_that_fits(client, cloud, no_job_threads):
    ids = ", ".join(str(n) for n in range(100000, 130000))
    body = re.sub(r"\s+", " ", client.post("/cota/jobs/preview", data={
        "devices": ids, "commands": "\n".join(["DAD76F4B"] * 8)}).text)
    assert "past the end of today" in body and "devices in progress at once" in body
    assert "Start job —" not in body                           # held until it fits the day

def test_a_refused_start_goes_back_to_the_form_as_it_was(client, cloud, no_job_threads):
    reply = _start_job(client, commands="DAD76F4B\nhello")
    assert reply.status_code == 200 and "fix these commands first" in reply.text
    assert "DAD76F4B\nhello</textarea>" in reply.text and 'action="/cota/jobs/preview"' in reply.text
    assert no_job_threads == []


def test_duplicate_and_rerun_fill_a_new_job(client, cloud, no_job_threads):
    _start_job(client)
    conn = db.connect()
    with conn:                          # as if 14906 answered everything and 786 did not
        conn.execute("UPDATE cota_campaign_result SET state = 'done' WHERE device_id = 14906")
        conn.execute("UPDATE cota_campaign_result SET state = 'failed' WHERE device_id = 786")
        conn.execute("UPDATE cota_campaign_device SET state = 'done'")
        conn.execute("UPDATE cota_campaign SET state = 'done'")
    page = client.get("/cota/jobs/1").text
    assert 'href="/cota/jobs/new?from=1"' in page and 'href="/cota/jobs/new?from=1&which=unfinished"' in page
    rerun = client.get("/cota/jobs/new?from=1&which=unfinished").text
    assert "786</textarea>" in rerun and "did not answer every command" in rerun
    copy = client.get("/cota/jobs/new?from=1").text
    assert "14906, 786</textarea>" in copy and "DAD76F4B\nDAD76C0A</textarea>" in copy
    assert cloud.sent == []                                           # nothing starts from here


def test_typed_commands_are_named_by_this_install(client, cota_env):
    out = client.post("/cota/commands/describe", data={"text": "DAD76F4B\nhello"}).json()
    assert [(l["line"], l["name"]) for l in out["lines"]] == [(1, "GET FTP_SETTINGS"), (2, "")]


def test_the_command_library_names_commands_on_every_page(client, cloud, no_job_threads):
    body = client.post("/cota/commands/parameters", data={"code": "6c0a", "name": "TIMERS"}).text
    assert "Parameter 6C0A is “TIMERS”." in body
    reply = client.post("/cota/commands/saved", data={"name": "Ignition timer 1 s",
                                                      "val1": "DBD76B82D531", "tags": "timers"},
                        follow_redirects=False)
    assert reply.status_code == 303
    library = re.sub(r"\s+", " ", client.get("/cota/commands").text)
    assert "Ignition timer 1 s" in library and "tag-chip" in library
    # An unnamed parameter shows by its code; the type has its own column.
    assert '<td class="param-cell"><span class="mono">6B82</span></td>' in library
    assert '<span class="pill pill-kind-set">SET</span>' in library
    preview = client.post("/cota/jobs/preview", data={"devices": "786",
                                                      "commands": "DAD76C0A\nDBD76B82D531"}).text
    assert "GET TIMERS" in preview and "Ignition timer 1 s" in preview
    assert 'data-val1="DBD76B82D531"' in client.get("/cota/jobs/new").text        # the picker
    assert 'data-library="single" data-val1="DBD76B82D531"' in client.get("/cota/console?device=786").text


def test_a_bad_saved_command_is_refused_and_kept_in_the_form(client, cota_env):
    body = client.post("/cota/commands/saved", data={"name": "Broken", "val1": "hello"}).text
    assert "Not saved" in body and 'value="hello"' in body and 'value="Broken"' in body


def test_a_saved_command_name_never_reaches_a_script_string(client, cota_env):
    client.post("/cota/commands/saved", data={"name": "x');alert(1);('", "val1": "DAD76F4B"})
    body = client.get("/cota/commands").text
    assert "onsubmit=" not in body
    assert 'data-confirm="Delete the saved command “{name}”? This cannot be undone."' in body


def test_saved_commands_can_be_edited_filtered_and_deleted(client, cota_env):
    client.post("/cota/commands/saved", data={"name": "FTP", "val1": "DAD76F4B", "tags": "read"})
    client.post("/cota/commands/saved", data={"name": "SOS off", "val1": "DDD76D66", "tags": "alarm"})
    cid = db.connect().execute("SELECT id FROM cota_saved_command WHERE name = 'FTP'").fetchone()[0]
    edit = client.get(f"/cota/commands?edit={cid}").text
    assert f'name="command_id" value="{cid}"' in edit and "Save changes" in edit
    only = client.get("/cota/commands?tag=alarm").text
    assert "SOS off" in only and ">FTP<" not in only
    client.post(f"/cota/commands/saved/{cid}/delete")
    assert db.connect().execute("SELECT COUNT(*) FROM cota_saved_command").fetchone()[0] == 1


def test_a_job_shows_tiles_and_pictures_instead_of_pills(client, cloud, no_job_threads):
    _start_job(client)
    body = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert 'class="count-pill' not in body.split('id="job-live"')[1]
    for label in ("Answered", "Waiting for devices", "Failed", "Expired", "Not sent yet"):
        assert f'<div class="tile-label">{label}</div>' in body
    assert 'class="stack stack-big"' in body
    # The time-to-answer and answered-on-attempt pictures were taken out at the user's request;
    # the same figures stay per command in the summary.
    assert "Answered on attempt" not in body and "First try" in body


def test_by_command_is_now_the_command_summary(client, cloud, no_job_threads):
    _start_job(client)
    body = client.get("/cota/jobs/1").text
    assert "Command summary" in body and "By command" not in body
    assert '<th class="n fit">In progress</th>' in body and "In flight" not in body     # the user's word, 10-10-2026
    assert body.count('<td class="outcome-col">') == 2              # one bar per command


def test_the_jobs_page_has_todays_dashboard_once_there_are_jobs(client, cloud, no_job_threads):
    empty = client.get("/cota").text
    assert "Jobs today" not in empty and "No jobs today" in empty
    _start_job(client)
    body = re.sub(r"\s+", " ", client.get("/cota").text)
    for label in ("Jobs today", "Devices reached", "Commands answered", "Time to answer", "Failed",
                  "Cloud calls"):
        assert f'<div class="tile-label">{label}</div>' in body
    assert "Outcomes per job" in body and 'class="job-bar-row" href="/cota/jobs/1"' in body
    assert "Answers per hour" in body and "No answers yet today." in body
    assert cloud.sent == [] and cloud.checks == []


def test_each_device_says_what_happens_next_and_when(client, cloud, no_job_threads):
    _start_job(client)
    body = re.sub(r"\s+", " ", client.get("/cota/jobs/1").text)
    assert re.search(r'class="next-what"[^>]*>Sends</span>', body) and ">due now</span>" in body
    assert 'id="job-clock" data-now="' in body and 'class="countdown" data-at="' in body
    conn = db.connect()
    with conn:
        conn.execute("UPDATE cota_campaign SET state = 'paused'")
    paused = client.get("/cota/jobs/1/status").json()["html"]
    assert '<span class="dim">paused</span>' in paused and "next-what" not in paused


def _source_state(body):
    """{source: (greyed, field disabled)} for the three device sources on the New job form."""
    out = {}
    for key, tag in (("group", "<select"), ("ids", "<textarea"), ("csv", '<input type="file"')):
        head = re.search(r'<div class="source-option([^"]*)" data-source="%s"[^>]*>' % key, body)
        field = re.search(re.escape(tag) + r"[^>]*>", body[head.end():]).group(0)
        out[key] = ("is-off" in head.group(1), " disabled" in field)
    return out


def _tab(body):
    """Which device-source tab the New job form opens on."""
    return re.search(r'data-tab="(\w+)" aria-selected="true"', body).group(1)


def test_a_new_job_takes_devices_from_one_source_at_a_time(client, cloud, no_job_threads):
    """Tabs (12-10-2026): only the chosen source is shown and sent — the others are hidden and
    disabled — and the page comes back from a preview on the source that was used."""
    blank = client.get("/cota/jobs/new").text
    assert _tab(blank) == "ids"                                     # no saved group yet
    assert _source_state(blank) == {"group": (True, True), "ids": (False, False), "csv": (True, True)}
    client.post("/cota/groups", data={"name": "Desk", "devices": "14906, 786"})
    gid = db.connect().execute("SELECT id FROM cota_group").fetchone()[0]
    assert _tab(client.get("/cota/jobs/new").text) == "group"       # groups first, once there are some
    by_group = client.post("/cota/jobs/preview", data={"group_id": str(gid), "commands": "DAD76F4B"}).text
    assert _tab(by_group) == "group"
    assert _source_state(by_group) == {"group": (False, False), "ids": (True, True), "csv": (True, True)}
    by_ids = client.post("/cota/jobs/preview", data={"devices": "786", "commands": "DAD76F4B"}).text
    assert _tab(by_ids) == "ids"
    assert _source_state(by_ids) == {"group": (True, True), "ids": (False, False), "csv": (True, True)}


# ─── Commands: search, paging, folds, sharing (2.0.1) ───────────────────────

def _save(client, name, val1, tags=""):
    return client.post("/cota/commands/saved", data={"name": name, "val1": val1, "tags": tags},
                       follow_redirects=False)


def test_the_library_is_searched_counted_and_numbered(client, cota_env):
    for n, v in (("Read FTP", "DAD76F4B"), ("SOS off", "DDD76D66"), ("Timers", "DAD76C0A")):
        _save(client, n, v)
    body = re.sub(r"\s+", " ", client.get("/cota/commands").text)
    assert 'name="q"' in body and re.search(r'>Saved commands</span> <span class="hint">3</span>', body)
    found = re.sub(r"\s+", " ", client.get("/cota/commands?q=sos").text)
    assert '<span class="hint">1 of 3</span>' in found and "SOS off" in found and "Read FTP" not in found
    assert re.search(r'<td class="n dim">1</td> <td class="name-cell"><strong[^>]*>SOS off', found)  # S.No.


def test_saved_commands_sort_on_every_header(client, cota_env):
    for n, v in (("b Read FTP", "DAD76F4B"), ("a SOS off", "DDD76D66"), ("c Timers", "DBD76C0AD531")):
        _save(client, n, v)
    names = lambda q: re.findall(r'<td class="name-cell"><strong[^>]*>([^<]+)</strong>', client.get("/cota/commands?" + q).text)
    assert names("") == ["b Read FTP", "c Timers", "a SOS off"]                 # GET, SET, CLR
    assert names("csort=name") == ["a SOS off", "b Read FTP", "c Timers"]
    assert names("csort=name&cdir=desc") == ["c Timers", "b Read FTP", "a SOS off"]
    page = client.get("/cota/commands?csort=name&cdir=desc").text
    assert 'aria-sort="descending"' in page and "csort=name&amp;cdir=desc" in page    # the pager keeps it


def test_the_library_pages_at_5_20_or_50(client, cota_env):
    for i in range(25):
        _save(client, f"Cmd {i:02d}", "DAD76F4B")
    flat = lambda path: re.sub(r"\s+", " ", client.get(path).text)
    first = flat("/cota/commands")
    assert "1–20 of <strong>25</strong>" in first and "Cmd 19" in first and "Cmd 20" not in first
    second = flat("/cota/commands?cpage=2")
    assert "21–25 of <strong>25</strong>" in second and '<td class="n dim">21</td>' in second
    assert "21–25 of <strong>25</strong>" in flat("/cota/commands?csize=5&cpage=5")
    assert "1–20 of <strong>25</strong>" in flat("/cota/commands?csize=7")      # not a size: 20


def test_one_dialog_adds_and_edits_and_parameter_names_fold_away(client, cota_env):
    """10-10-2026: saved commands are the page; one Add command dialog, not two forms; parameter
    names folded at the bottom, renamed or removed there."""
    body = client.get("/cota/commands").text
    assert body.count("<button type=\"button\" class=\"btn-primary btn-icon\" data-add-command>") == 1 and 'id="param-form"' not in body
    assert re.search(r'<dialog class="form-dialog" id="cmd-dialog"[^>]*\s>', body)       # closed
    assert "data-autoopen" not in body.split('id="cmd-dialog"')[1].split(">")[0]
    assert re.search(r'<details class="card fold-card" id="parameters"\s*>', body)    # folded
    assert body.index('id="saved"') < body.index('id="parameters"')
    _save(client, "Read FTP", "DAD76F4B")
    cid = db.connect().execute("SELECT id FROM cota_saved_command").fetchone()[0]
    editing = client.get(f"/cota/commands?edit={cid}").text
    assert re.search(r'id="cmd-dialog"[^>]*data-autoopen', editing)                 # open, filled in
    assert re.search(r'<details class="card fold-card" id="parameters"\s*open>',
                     client.get("/cota/commands?names=1").text)


def test_a_command_names_its_parameter_in_the_same_save(client, cota_env):
    out = client.post("/cota/commands/describe", data={"text": "DBD76B82D531"}).json()["lines"][0]
    assert (out["parameter"], out["parameter_name"], out["reads_as"]) == ("6B82", "", "SET 6B82")
    reply = client.post("/cota/commands/saved", data={"name": "Ignition timer 1 s", "val1": "DBD76B82D531",
                                                      "parameter_name": "IGN_TIMER"}, follow_redirects=False)
    assert reply.status_code == 303
    conn = db.connect()
    assert conn.execute("SELECT name FROM cota_parameter WHERE code = '6B82'").fetchone()[0] == "IGN_TIMER"
    named = client.post("/cota/commands/describe", data={"text": "DAD76B82"}).json()["lines"][0]
    assert named["parameter_name"] == "IGN_TIMER" and named["saved_as"] == ""
    again = client.post("/cota/commands/describe", data={"text": "DBD76B82D531"}).json()["lines"][0]
    assert again["saved_as"] == "Ignition timer 1 s"
    # A blank parameter name names nothing; a refused save keeps what was typed.
    client.post("/cota/commands/saved", data={"name": "Get 6C0A", "val1": "DAD76C0A", "parameter_name": " "})
    assert conn.execute("SELECT COUNT(*) FROM cota_parameter WHERE code = '6C0A'").fetchone()[0] == 0
    refused = client.post("/cota/commands/saved", data={"name": "", "val1": "DAD76D00",
                                                        "parameter_name": "X1"}).text
    assert 'value="X1"' in refused


def test_the_commands_page_carries_no_help_text(client, cota_env):
    _save(client, "Read FTP", "DAD76F4B")
    visible = re.sub(r'data-tip="[^"]*"', "", client.get("/cota/commands").text)
    for phrase in ("used by", "every page and export", "Your names for", " · was "):
        assert phrase not in visible, phrase


def test_saving_a_command_already_saved_under_another_name_says_so(client, cota_env):
    _save(client, "SOS off", "DDD76D66")
    reply = _save(client, "Clear SOS", "DDD76D66")
    body = client.get(reply.headers["location"]).text
    assert "Saved “Clear SOS”. The same command is also saved as “SOS off”." in body


def test_the_library_exports_and_imports_through_a_preview(client, cota_env):
    _save(client, "Read FTP", "DAD76F4B", "read")
    client.post("/cota/commands/parameters", data={"code": "6C0A", "name": "TIMERS"})
    exported = client.get("/cota/commands/export.csv")
    assert exported.status_code == 200 and "cota_commands" in exported.headers["content-disposition"]
    text = exported.content.decode("utf-8-sig")
    assert text.splitlines()[0] == "kind,name,command,tags,note,wait_after" and "command,Read FTP,DAD76F4B,read," in text
    assert client.get("/cota/commands/template.csv").content.decode("utf-8-sig").startswith("kind,name,command")

    upload = ("kind,name,command,tags,note\ncommand,SOS off,DDD76D66,alarm,\n"
              "command,Read FTP,DAD76F4B,read,\n")
    preview = client.post("/cota/commands/import",
                          files={"file": ("lib.csv", upload.encode(), "text/csv")}).text
    assert ">Import preview</span>" in preview and "Apply — 1 new, 0 updated" in preview
    assert db.connect().execute("SELECT COUNT(*) FROM cota_saved_command").fetchone()[0] == 1
    content = re.search(r'<textarea name="content" hidden>(.*?)</textarea>', preview, re.S).group(1)
    done = client.post("/cota/commands/import/apply", data={"content": html.unescape(content)},
                       follow_redirects=False)
    assert done.headers["location"] == "/cota/commands?imported=1,0,1,0"
    assert "Imported: 1 new, 0 updated, 1 already the same." in client.get(done.headers["location"]).text
    assert db.connect().execute("SELECT COUNT(*) FROM cota_saved_command").fetchone()[0] == 2


def test_an_import_without_a_file_or_the_columns_says_why(client, cota_env):
    assert "Choose a file to import" in client.post("/cota/commands/import", data={}).text
    bad = client.post("/cota/commands/import", files={"file": ("x.csv", b"a,b\n1,2\n", "text/csv")}).text
    assert "kind, name and command" in bad


def test_a_saved_command_starts_a_new_job(client, cota_env):
    _save(client, "Read FTP", "DAD76F4B")
    assert 'href="/cota/jobs/new?commands=DAD76F4B"' in client.get("/cota/commands").text
    new = client.get("/cota/jobs/new?commands=DAD76F4B").text
    assert re.search(r'id="job-commands"[^>]*>DAD76F4B</textarea>', new)


def test_a_long_library_gets_a_filter_in_the_picker(client, cota_env):
    for i in range(9):
        _save(client, f"Cmd {i}", "DAD76F4B")
    assert "data-library-filter" in client.get("/cota/jobs/new").text



def test_the_stylesheet_url_changes_when_the_stylesheet_does(client, monkeypatch, tmp_path):
    """Versioned by the app version alone, the URL stayed the same through a day of style
    changes and the browser kept its old copy: new parts of pages came out unstyled."""
    from ota_analytics import api

    first = api.static_version()
    css = tmp_path / "static" / "app.css"
    css.parent.mkdir()
    css.write_text("body {}", encoding="utf-8")
    monkeypatch.setattr(api, "WEB", tmp_path)
    import os
    os.utime(css, (1_800_000_000, 1_800_000_000))
    assert api.static_version().endswith(".1800000000") and api.static_version() != first
    monkeypatch.undo()
    assert f'app.css?v={api.static_version()}"' in client.get("/cota").text



def test_every_delete_and_cancel_asks_first_and_is_red(client, cloud, no_job_threads):
    """The user's standard: destructive actions are red, and each opens a confirmation."""
    client.post("/cota/groups", data={"name": "Desk", "devices": "14906, 786"})
    client.post("/cota/commands/saved", data={"name": "Read FTP", "val1": "DAD76F4B"})
    client.post("/cota/commands/parameters", data={"code": "6C0A", "name": "TIMERS"})
    client.post("/cota/jobs/start", data={"device_ids": "786", "commands": "DAD76F4B"})
    for path in ("/cota/devices", "/cota/commands", "/cota/jobs/1"):
        body = client.get(path).text
        destructive = re.findall(r'<button[^>]*>(?:(?!</button>).)*?<span>(Delete|Remove|Reset|Cancel)</span>',
                                 body, re.S)
        assert destructive, path
        for button in re.findall(r'<button[^>]*danger[^>]*>', body):
            assert "danger" in button
        assert body.count("data-confirm=") >= len(destructive), path
    assert 'id="confirm-dialog"' in client.get("/cota/commands").text            # one dialog, in base


# ─── one design through the application (12-10-2026) ─────────────────────────

def _devices_total(body: str) -> int:
    hint = re.search(r'<span class="hint">([\d,]+)', re.sub(r"\s+", " ", body))
    return int(hint.group(1).replace(",", ""))


def test_a_list_filter_takes_several_values_and_ors_them(client):
    online = _devices_total(client.get("/devices?status=Online").text)
    offline = _devices_total(client.get("/devices?status=Offline").text)
    both = _devices_total(client.get("/devices?status=Online&status=Offline").text)
    assert online and offline and both == online + offline


def test_a_list_filter_offers_only_values_some_device_has(client):
    """"0 means it does not exist" — the user. A status no device has is not offered."""
    from ota_analytics import db, fota_devices, metrics

    conn = db.connect()
    sid = metrics.latest_snapshot_id(conn)
    for key, values in fota_devices.options(conn, sid)["choices"].items():
        assert all(n > 0 for _, _, n in values), key


def test_a_pasted_imei_list_matches_exactly_those(client):
    body = client.get("/devices?size=25").text
    imeis = re.findall(r'class="mono">(\d+)<', body)[:3]
    assert len(imeis) == 3
    import urllib.parse
    page = client.get("/devices?q=" + urllib.parse.quote("\n".join(imeis))).text
    assert sorted(re.findall(r'class="mono">(\d+)<', page)) == sorted(imeis)


def test_the_last_change_and_last_seen_filters_are_buckets_that_add_up(client):
    from ota_analytics import db, fota_devices, metrics

    conn = db.connect()
    sid = metrics.latest_snapshot_id(conn)
    total = fota_devices.rows(conn, sid, {})[1]
    for key in ("seen", "changed"):
        counts = {v: n for v, _, n in fota_devices.options(conn, sid)["choices"][key]}
        assert sum(counts.values()) == total, key
        for value, n in counts.items():
            assert fota_devices.rows(conn, sid, {key: [value]})[1] == n, (key, value)


def test_old_last_update_links_still_filter(client):
    """?changed=24h came from the old filter bar and may be bookmarked."""
    assert client.get("/devices?changed=24h").status_code == 200


def test_the_devices_list_has_the_managed_list_head(client):
    flat = re.sub(r"\s+", " ", client.get("/devices").text)
    assert 'class="table-wrap grid-scroll" id="grid" data-fit' in flat
    assert "Clear filters" not in flat                 # only while filtered
    filtered = re.sub(r"\s+", " ", client.get("/devices?status=Online").text)
    assert "Clear filters" in filtered
    assert re.search(r'<span class="hint">[\d,]+ of [\d,]+</span>', filtered)
    for fmt in ("xlsx", "csv", "txt"):
        assert f"/devices/export?" in filtered and f"format={fmt}" in filtered


def test_the_export_follows_the_list_filters(client):
    import csv
    import io

    body = client.get("/devices/export?status=Online&status=Offline&format=csv").text
    rows = list(csv.DictReader(io.StringIO(body.lstrip("\ufeff"))))
    page_total = _devices_total(client.get("/devices?status=Online&status=Offline").text)
    assert len(rows) == page_total


def test_fota_tiles_lead_with_the_share(client):
    """The user's order on every tile: the percentage the focus, the count beside it, smaller."""
    for path in ("/", "/pending", "/changes?window=all"):
        assert 'class="tile-count"' in client.get(path).text, path


def test_clearing_the_error_log_is_red_and_confirmed(client):
    from ota_analytics import db, errors

    errors.record("page", RuntimeError("boom"), path="/x", conn=db.connect())
    flat = re.sub(r"\s+", " ", client.get("/errors").text)
    form = flat[flat.index('action="/errors/clear"') - 200:]
    form = form[:form.index("</form>")]
    assert "data-confirm=" in form and "danger" in form
